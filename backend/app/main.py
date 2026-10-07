import os, json, datetime as dt, io, re, unicodedata, asyncio, logging, base64, hashlib, hmac, csv, math
import html as html_lib
import smtplib, ssl
from email.message import EmailMessage
from email.utils import formataddr
from pathlib import Path
from urllib.parse import quote
from zoneinfo import ZoneInfo
from alembic import command
from alembic.config import Config
from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File, Form, Request
from fastapi.responses import Response, JSONResponse
import jwt
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel
from pydantic import Field
from sqlalchemy import select, func, inspect, case, delete, or_, and_
from sqlalchemy.orm import Session, selectinload
from sqlalchemy.exc import IntegrityError
from .db import Base, engine, get_db, SessionLocal
from .models import *
from .security import hash_pw, check_pw, make_token, make_password_reset_token, require, current_user
from cryptography.fernet import Fernet, InvalidToken

MACHINE_STATES = {"OPERATIVA","OPERATIVA_CON_OBSERVACION","EN_MANTENIMIENTO","FUERA_DE_SERVICIO","PENDIENTE_DE_REPUESTO","RETIRADA"}
WRITERS = ("JEFE","SUPERVISOR","TECNICO")
TICKET_SLA_HOURS = {"CRÍTICA": 5, "ALTA": 24, "NORMAL": 120, "BAJA": 168}

app = FastAPI(title="Gestión Técnica API", version="0.1.0")
logger = logging.getLogger(__name__)

@app.middleware("http")
async def log_unhandled_request_errors(request: Request, call_next):
    try:
        return await call_next(request)
    except Exception:
        logger.exception("Error no controlado en %s %s", request.method, request.url.path)
        raise

@app.exception_handler(Exception)
async def log_unhandled_api_errors(request: Request, exc: Exception):
    logger.error("Excepción no controlada en %s %s", request.method, request.url.path,
                 exc_info=(type(exc), exc, exc.__traceback__))
    return JSONResponse(status_code=500, content={"detail": "Internal Server Error"})

def purge_expired_ticket_attachments():
    cutoff = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=7)
    latest_closed = (select(TicketStatusHistory.ticket_id.label("ticket_id"),
                            func.max(TicketStatusHistory.at).label("closed_at"))
                     .where(TicketStatusHistory.new_status == "CERRADO")
                     .group_by(TicketStatusHistory.ticket_id).subquery())
    expired_tickets = (select(Ticket.id).join(latest_closed, latest_closed.c.ticket_id == Ticket.id)
                       .where(Ticket.status == "CERRADO", latest_closed.c.closed_at <= cutoff))
    with SessionLocal() as db:
        result = db.execute(delete(TicketAttachment).where(TicketAttachment.ticket_id.in_(expired_tickets)))
        db.commit()
        return result.rowcount or 0

async def ticket_attachment_retention_worker():
    while True:
        try:
            removed = await asyncio.to_thread(purge_expired_ticket_attachments)
            if removed:
                logger.info("Eliminados %s adjuntos de tickets cerrados hace más de 7 días", removed)
        except Exception:
            logger.exception("No se pudieron limpiar los adjuntos vencidos de tickets")
        await asyncio.sleep(6 * 60 * 60)

def render_floorplan(pdf_data: bytes) -> bytes:
    import fitz
    try:
        document = fitz.open(stream=pdf_data, filetype="pdf")
        if document.page_count < 1: raise ValueError("El PDF no contiene páginas")
        page = document[0]
        scale = min(150 / 72, 3500 / max(page.rect.width, page.rect.height))
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        image = pixmap.tobytes("png")
        document.close()
        return image
    except Exception as exc:
        raise ValueError(f"No se pudo procesar la primera página del PDF: {exc}") from exc

LEGACY_TABLES = set(Base.metadata.tables) - {"ticket_comments", "floor_plans", "ticket_attachments", "email_settings", "preventive_maintenance_plans", "preventive_maintenance_logs", "intervention_follow_ups"}

def migrate_schema():
    config = Config(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
    existing_tables = set(inspect(engine).get_table_names())
    if "alembic_version" not in existing_tables:
        legacy_tables = existing_tables & LEGACY_TABLES
        if legacy_tables:
            missing = LEGACY_TABLES - existing_tables
            if missing:
                raise RuntimeError(
                    "La base contiene un esquema antiguo incompleto; no se aplicaron migraciones. "
                    f"Faltan tablas: {', '.join(sorted(missing))}"
                )
            inspector = inspect(engine)
            missing_columns = []
            migrated_columns = {"interventions": {"island", "jira_number", "pending"},
                                "users": {"technician_id", "avatar_data", "avatar_content_type"},
                                "technicians": {"avatar_data", "avatar_content_type"},
                                "tickets": {"sla_due_at", "sla_overdue_notified_at", "sla_approaching_notified_at"},
                                "machines": {"position_x", "position_y"}}
            for table_name in sorted(LEGACY_TABLES):
                present = {column["name"] for column in inspector.get_columns(table_name)}
                expected = {column.name for column in Base.metadata.tables[table_name].columns}
                expected -= migrated_columns.get(table_name, set())
                absent = expected - present
                if absent:
                    missing_columns.append(f"{table_name} ({', '.join(sorted(absent))})")
            if missing_columns:
                raise RuntimeError(
                    "La base contiene un esquema antiguo incompleto; no se aplicaron migraciones. "
                    f"Faltan columnas: {'; '.join(missing_columns)}"
                )
            # Las instalaciones previas usaban create_all; se marca el esquema existente
            # para conservar sus datos y aplicar solo las revisiones posteriores.
            command.stamp(config, "0001_initial")
    command.upgrade(config, "head")

@app.on_event("startup")
def startup():
    migrate_schema()
    with SessionLocal() as db:
        if not db.scalar(select(User).limit(1)):
            db.add(User(username=os.environ["ADMIN_USER"], full_name="Administrador",
                        password_hash=hash_pw(os.environ["ADMIN_PASSWORD"]), role="ADMIN"))
            db.add_all([Area(name=n) for n in ["Slots","Bunker","Audiovisual","Sistemas","Otros"]])
            db.commit()
        if not db.get(FloorPlan, 1):
            source = Path(__file__).resolve().parents[1] / "assets" / "layout-sala-inicial.pdf"
            if source.exists():
                pdf_data = source.read_bytes()
                db.add(FloorPlan(id=1, filename="01-09-2026_Layout_Sala_OCR_09-2026_V21.1.pdf",
                                 pdf_data=pdf_data, image_data=render_floorplan(pdf_data), uploaded_by=None))
                db.commit()
    app.state.ticket_attachment_retention_task = asyncio.create_task(ticket_attachment_retention_worker())
    app.state.sla_overdue_notification_task = asyncio.create_task(sla_overdue_notification_worker())

@app.on_event("shutdown")
async def stop_ticket_attachment_retention_worker():
    tasks = [getattr(app.state, "ticket_attachment_retention_task", None),
             getattr(app.state, "sla_overdue_notification_task", None)]
    for task in tasks:
        if not task: continue
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass

def audit(db, user, action, entity, old=None, new=None):
    db.add(AuditLog(user_id=user.id, action=action, entity=entity,
                    old_value=json.dumps(old, default=str) if old is not None else None,
                    new_value=json.dumps(new, default=str) if new is not None else None))

@app.get("/admin/audit")
def list_audit_logs(page: int = Query(1, ge=1), size: int = Query(50, ge=1, le=200),
                   q: str | None = None, db: Session = Depends(get_db), _=Depends(require("ADMIN"))):
    query = select(AuditLog, User).outerjoin(User, User.id == AuditLog.user_id)
    if q:
        term = f"%{q.strip()}%"
        query = query.where((AuditLog.action.ilike(term)) | (AuditLog.entity.ilike(term)) |
                            (User.username.ilike(term)) | (User.full_name.ilike(term)))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.execute(query.order_by(AuditLog.at.desc(), AuditLog.id.desc())
                      .offset((page - 1) * size).limit(size)).all()
    return {"total": total, "items": [{"id": log.id, "at": log.at,
             "user": user.full_name or user.username if user else "Sistema",
             "action": log.action, "entity": log.entity,
             "old_value": log.old_value, "new_value": log.new_value} for log, user in rows]}

def send_system_email(to: list[str], subject: str, body: str):
    config = active_email_settings()
    if not (config["smtp_host"] and config["smtp_user"] and config["smtp_password"] and to):
        logger.warning("Correo no enviado: falta configuración SMTP o destinatarios")
        return False
    message = EmailMessage()
    message["Subject"] = subject.replace("\r", " ").replace("\n", " ")
    message["From"] = formataddr((config["smtp_from_name"], config["smtp_from"] or config["smtp_user"]))
    message["To"] = ", ".join(to)
    message.set_content(body)
    return send_email_message(message, config)

def send_email_message(message: EmailMessage, config: dict | None = None):
    config = config or active_email_settings()
    try:
        port = int(config["smtp_port"])
    except ValueError:
        logger.error("SMTP_PORT debe ser un número")
        return False
    host, username, password = config["smtp_host"], config["smtp_user"], config["smtp_password"]
    use_tls = bool(config["smtp_use_tls"])
    try:
        if use_tls and port == 465:
            with smtplib.SMTP_SSL(host, port, timeout=20, context=ssl.create_default_context()) as smtp:
                smtp.login(username, password); smtp.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=20) as smtp:
                if use_tls: smtp.starttls(context=ssl.create_default_context())
                smtp.login(username, password); smtp.send_message(message)
        return True
    except (OSError, smtplib.SMTPException):
        logger.exception("No se pudo enviar correo del sistema")
        return False

def email_cipher():
    secret = os.environ["JWT_SECRET"].encode()
    key = base64.urlsafe_b64encode(hmac.new(secret, b"gestion-tecnica/email-settings/v1", hashlib.sha256).digest())
    return Fernet(key)

def active_email_settings():
    with SessionLocal() as db:
        saved = db.get(EmailSettings, 1)
        if saved:
            encrypted = saved.smtp_password_encrypted
            password = os.getenv("SMTP_PASSWORD", "")
            if encrypted:
                try:
                    password = email_cipher().decrypt(encrypted.encode()).decode()
                except InvalidToken:
                    logger.error("No se pudo descifrar la contraseña SMTP guardada; revisa JWT_SECRET")
                    password = ""
            return {"smtp_host": saved.smtp_host or "", "smtp_port": saved.smtp_port,
                    "smtp_user": saved.smtp_user or "", "smtp_password": password,
                    "smtp_from": saved.smtp_from or "", "smtp_from_name": saved.smtp_from_name,
                    "smtp_use_tls": saved.smtp_use_tls}
    return {"smtp_host": os.getenv("SMTP_HOST", "").strip(),
            "smtp_port": int(os.getenv("SMTP_PORT", "587")),
            "smtp_user": os.getenv("SMTP_USER", "").strip(),
            "smtp_password": os.getenv("SMTP_PASSWORD", ""),
            "smtp_from": os.getenv("SMTP_FROM", "").strip(),
            "smtp_from_name": os.getenv("SMTP_FROM_NAME", "Gestión Técnica"),
            "smtp_use_tls": os.getenv("SMTP_USE_TLS", "true").strip().lower() in {"1", "true", "yes", "si"}}

def email_recipient_settings():
    with SessionLocal() as db:
        saved = db.get(EmailSettings, 1)
        if saved:
            return {"NOTIFICATION_EMAIL_TO": saved.notification_email_to,
                    "REPORT_EMAIL_TO": saved.report_email_to,
                    "REPORT_EMAIL_CC": saved.report_email_cc}
    return {key: os.getenv(key, "") for key in
            ("NOTIFICATION_EMAIL_TO", "REPORT_EMAIL_TO", "REPORT_EMAIL_CC")}

def email_settings_dict(db: Session):
    saved = db.get(EmailSettings, 1)
    config = active_email_settings()
    recipients = email_recipient_settings()
    return {"smtp_host": config["smtp_host"], "smtp_port": config["smtp_port"],
            "smtp_user": config["smtp_user"], "smtp_from": config["smtp_from"],
            "smtp_from_name": config["smtp_from_name"], "smtp_use_tls": config["smtp_use_tls"],
            "smtp_password_set": bool(config["smtp_password"]),
            **recipients, "updated_at": saved.updated_at if saved else None,
            "updated_by": saved.updated_by if saved else None}

def send_ticket_notification(subject: str, body: str):
    recipients = report_recipients("NOTIFICATION_EMAIL_TO") or report_recipients("REPORT_EMAIL_TO")
    return send_system_email(recipients, subject, body)

def notify_overdue_sla_tickets():
    now_ = dt.datetime.now(dt.timezone.utc)
    with SessionLocal() as db:
        overdue = db.scalars(select(Ticket).where(
            Ticket.status.notin_({"RESUELTO", "CERRADO", "CANCELADO"}),
            Ticket.sla_due_at.is_not(None), Ticket.sla_due_at <= now_,
            Ticket.sla_overdue_notified_at.is_(None))
            .order_by(Ticket.sla_due_at).limit(100).with_for_update(skip_locked=True)).all()
        sent = 0
        for ticket in overdue:
            subject = f"SLA vencido · Ticket #{ticket.id} · {ticket.priority}"
            body = (f"El ticket #{ticket.id} superó su plazo SLA.\n"
                    f"Tarea: {ticket.task}\nPrioridad: {ticket.priority}\n"
                    f"Vencimiento: {ticket.sla_due_at:%d/%m/%Y %H:%M UTC}\n"
                    f"Técnico: {ticket.technician or 'Sin asignar'}\n"
                    "Ingresa a Gestión Técnica para revisarlo.")
            if send_ticket_notification(subject, body):
                ticket.sla_overdue_notified_at = now_
                db.add(AuditLog(user_id=None, action="NOTIFICACION_SLA_VENCIDO",
                                entity=f"ticket:{ticket.id}",
                                new_value=json.dumps({"sent_to_team": True, "sla_due_at": ticket.sla_due_at.isoformat()})))
                sent += 1
        db.commit()
        return sent

async def sla_overdue_notification_worker():
    while True:
        try:
            sent = await asyncio.to_thread(notify_sla_tickets)
            maintenance_sent = await asyncio.to_thread(notify_due_maintenance_plans)
            if sent: logger.info("Enviados %s avisos de SLA", sent)
            if maintenance_sent: logger.info("Enviados %s avisos de mantenimiento preventivo", maintenance_sent)
        except Exception:
            logger.exception("No se pudieron enviar los avisos de SLA vencido")
        await asyncio.sleep(15 * 60)

def notify_sla_tickets():
    now_ = dt.datetime.now(dt.timezone.utc)
    with SessionLocal() as db:
        approaching = db.scalars(select(Ticket).where(
            Ticket.status.notin_({"RESUELTO", "CERRADO", "CANCELADO"}),
            Ticket.sla_due_at > now_, Ticket.sla_due_at <= now_ + dt.timedelta(hours=24),
            Ticket.sla_approaching_notified_at.is_(None)).limit(100).with_for_update(skip_locked=True)).all()
        sent = 0
        for ticket in approaching:
            if send_ticket_notification(f"SLA próximo a vencer · Ticket #{ticket.id}",
                f"El ticket #{ticket.id} vencerá dentro de las próximas 24 horas.\nTarea: {ticket.task}\nPrioridad: {ticket.priority}\nVencimiento: {ticket.sla_due_at:%d/%m/%Y %H:%M UTC}\nTécnico: {ticket.technician or 'Sin asignar'}"):
                ticket.sla_approaching_notified_at = now_
                db.add(AuditLog(action="NOTIFICACION_SLA_PROXIMO", entity=f"ticket:{ticket.id}",
                    new_value=json.dumps({"sla_due_at": ticket.sla_due_at.isoformat()})))
                sent += 1
        db.commit()
    return sent + notify_overdue_sla_tickets()

def notify_due_maintenance_plans():
    now_ = dt.datetime.now(dt.timezone.utc)
    with SessionLocal() as db:
        plans = db.scalars(select(PreventiveMaintenancePlan).where(
            PreventiveMaintenancePlan.active.is_(True),
            PreventiveMaintenancePlan.next_due_at <= now_ + dt.timedelta(days=3),
            PreventiveMaintenancePlan.notified_for_due_at.is_distinct_from(PreventiveMaintenancePlan.next_due_at))
            .order_by(PreventiveMaintenancePlan.next_due_at).limit(100).with_for_update(skip_locked=True)).all()
        sent = 0
        for plan in plans:
            due = plan.next_due_at
            if send_ticket_notification(f"Mantenimiento preventivo · Máquina {plan.machine.number}",
                f"Plan: {plan.title}\nMáquina: {plan.machine.number}\nVencimiento: {due:%d/%m/%Y}\nDescripción: {plan.description or '—'}"):
                plan.notified_for_due_at = due
                db.add(AuditLog(action="NOTIFICACION_MANTENIMIENTO", entity=f"maintenance:{plan.id}",
                    new_value=json.dumps({"due_at": due.isoformat()})))
                sent += 1
        db.commit()
        return sent

@app.get("/health")
def health(): return {"ok": True}

@app.post("/auth/login")
def login(f: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    login_name = f.username.strip()
    u = db.scalar(select(User).where(User.username == login_name))
    if not u:
        u = db.scalar(select(User).where(func.lower(User.username) == login_name.lower()))
    if not u or not check_pw(f.password, u.password_hash):
        raise HTTPException(401, "Credenciales inválidas")
    return {"access_token": make_token(u), "token_type": "bearer", "role": u.role}

class UserRegistration(BaseModel):
    email: str
    full_name: str
    password: str

@app.post("/auth/register", status_code=201)
def register_user(payload: UserRegistration, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    full_name = payload.full_name.strip()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
        raise HTTPException(422, "Ingresa un correo electrónico válido")
    if len(email) > 50:
        raise HTTPException(422, "El correo no puede superar 50 caracteres")
    if not full_name or len(full_name) > 120:
        raise HTTPException(422, "Ingresa tu nombre completo (máximo 120 caracteres)")
    if len(payload.password) < 8:
        raise HTTPException(422, "La contraseña debe tener al menos 8 caracteres")
    if len(payload.password.encode()) > 72:
        raise HTTPException(422, "La contraseña no puede superar 72 bytes")
    if db.scalar(select(User.id).where(func.lower(User.username) == email)):
        raise HTTPException(409, "Ya existe una cuenta con ese correo")
    user = User(username=email, full_name=full_name, password_hash=hash_pw(payload.password), role="USUARIO")
    db.add(user)
    db.commit()
    return {"id": user.id, "email": user.username, "full_name": user.full_name, "role": user.role}

@app.get("/auth/me")
def auth_me(u: User = Depends(current_user)):
    return {"id": u.id, "email": u.username, "full_name": u.full_name, "role": u.role,
            "technician_id": u.technician_id,
            "has_avatar": bool(u.avatar_data),
            "technician_name": (f"{u.technician.first_name} {u.technician.last_name}" if u.technician else None)}

@app.get("/auth/me/avatar")
def get_my_avatar(u: User = Depends(current_user)):
    if not u.avatar_data or not u.avatar_content_type:
        raise HTTPException(404, "No has cargado una foto de perfil")
    return Response(content=u.avatar_data, media_type=u.avatar_content_type,
                    headers={"Cache-Control": "private, no-store"})

async def read_avatar_image(file: UploadFile):
    allowed_signatures = {
        "image/jpeg": (b"\xff\xd8\xff",),
        "image/png": (b"\x89PNG\r\n\x1a\n",),
        "image/webp": (b"RIFF",),
    }
    content_type = (file.content_type or "").lower()
    if content_type not in allowed_signatures:
        raise HTTPException(415, "Usa una imagen JPG, PNG o WebP")
    image = await file.read(2 * 1024 * 1024 + 1)
    if len(image) > 2 * 1024 * 1024:
        raise HTTPException(413, "La imagen supera el límite de 2 MB")
    valid_signature = any(image.startswith(signature) for signature in allowed_signatures[content_type])
    if content_type == "image/webp":
        valid_signature = valid_signature and image[8:12] == b"WEBP"
    if not valid_signature:
        raise HTTPException(415, "El archivo no coincide con el formato de imagen indicado")
    return image, content_type

@app.put("/auth/me/avatar")
async def update_my_avatar(file: UploadFile = File(...), db: Session = Depends(get_db),
                          u: User = Depends(current_user)):
    image, content_type = await read_avatar_image(file)
    u.avatar_data = image
    u.avatar_content_type = content_type
    audit(db, u, "ACTUALIZAR_AVATAR", f"user:{u.id}", None,
          {"content_type": content_type, "size": len(image)})
    db.commit()
    return {"has_avatar": True}

@app.delete("/auth/me/avatar")
def delete_my_avatar(db: Session = Depends(get_db), u: User = Depends(current_user)):
    u.avatar_data = None
    u.avatar_content_type = None
    audit(db, u, "ELIMINAR_AVATAR", f"user:{u.id}")
    db.commit()
    return {"has_avatar": False}

class PasswordRecoveryIn(BaseModel):
    email: str

class PasswordResetIn(BaseModel):
    token: str
    password: str

@app.post("/auth/password-recovery")
def password_recovery(payload: PasswordRecoveryIn, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    user = db.scalar(select(User).where(func.lower(User.username) == email))
    if user:
        reset_url = os.getenv("FRONTEND_URL", "http://localhost:8080").rstrip("/") + "/?reset=" + quote(make_password_reset_token(user))
        send_system_email([user.username], "Recuperación de contraseña · Gestión Técnica",
                          f"Hola {user.full_name},\n\nPara cambiar tu contraseña, abre este enlace (válido por 30 minutos):\n{reset_url}\n\nSi no solicitaste este cambio, ignora este correo.")
    return {"message": "Si el correo está registrado, recibirás instrucciones para recuperar el acceso."}

@app.post("/auth/password-reset")
def password_reset(payload: PasswordResetIn, db: Session = Depends(get_db)):
    if len(payload.password) < 8 or len(payload.password.encode()) > 72:
        raise HTTPException(422, "La contraseña debe tener entre 8 y 72 bytes")
    try:
        claims = jwt.decode(payload.token, os.environ["JWT_SECRET"], algorithms=["HS256"])
        if claims.get("purpose") != "password_reset": raise ValueError("purpose")
        user = db.get(User, int(claims["sub"]))
    except Exception:
        raise HTTPException(400, "El enlace de recuperación es inválido o venció")
    if not user or claims.get("password_hash") != user.password_hash:
        raise HTTPException(400, "El enlace de recuperación es inválido, ya fue usado o venció")
    user.password_hash = hash_pw(payload.password)
    audit(db, user, "CAMBIAR_CONTRASENA_RECUPERACION", f"user:{user.id}")
    db.commit()
    return {"message": "Contraseña actualizada. Ya puedes iniciar sesión."}

SYSTEM_ROLES = {"ADMIN", "JEFE", "SUPERVISOR", "TECNICO", "CONSULTA", "USUARIO"}

class SystemUserIn(BaseModel):
    username: str
    full_name: str
    role: str
    password: str | None = None
    technician_id: int | None = None

def system_user_dict(user: User):
    return {"id": user.id, "username": user.username, "full_name": user.full_name,
            "role": user.role, "technician_id": user.technician_id,
            "technician_name": (f"{user.technician.first_name} {user.technician.last_name}"
                                if user.technician else None)}

def validate_system_user(data: SystemUserIn, db: Session, user_id: int | None = None):
    username = data.username.strip().lower()
    full_name = data.full_name.strip()
    is_email = bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", username))
    existing_user = db.get(User, user_id) if user_id else None
    legacy_admin_name = bool(existing_user and existing_user.role == "ADMIN" and existing_user.username.lower() == username)
    if not is_email and not legacy_admin_name:
        raise HTTPException(422, "Ingresa un correo electrónico válido")
    if len(username) > 50: raise HTTPException(422, "El correo no puede superar 50 caracteres")
    if not full_name or len(full_name) > 120:
        raise HTTPException(422, "El nombre es obligatorio y no puede superar 120 caracteres")
    if data.role not in SYSTEM_ROLES: raise HTTPException(422, "Rol de usuario inválido")
    if data.password is not None and (len(data.password) < 8 or len(data.password.encode()) > 72):
        raise HTTPException(422, "La contraseña debe tener entre 8 y 72 bytes")
    if data.role == "TECNICO" and data.technician_id is None:
        raise HTTPException(422, "Asocia una ficha de técnico a esta cuenta")
    if data.role != "TECNICO" and data.technician_id is not None:
        raise HTTPException(422, "Solo las cuentas con rol Técnico pueden vincular una ficha")
    duplicate = db.scalar(select(User.id).where(func.lower(User.username) == username,
                                                User.id != (user_id or -1)))
    if duplicate: raise HTTPException(409, "Ya existe una cuenta con ese correo")
    technician = db.get(Technician, data.technician_id) if data.technician_id else None
    if data.technician_id and not technician: raise HTTPException(404, "La ficha de técnico no existe")
    if technician and technician.status == "INACTIVO":
        existing_user = db.get(User, user_id) if user_id else None
        if not existing_user or existing_user.technician_id != technician.id:
            raise HTTPException(422, "No puedes vincular un técnico inactivo")
    if technician:
        linked = db.scalar(select(User.id).where(User.technician_id == technician.id,
                                                 User.id != (user_id or -1)))
        if linked: raise HTTPException(409, "La ficha ya está vinculada a otra cuenta")
        username_owner = db.scalar(select(Technician.id).where(Technician.username == username,
                                                              Technician.id != technician.id))
        if username_owner: raise HTTPException(409, "El correo ya está asignado a otro técnico")
    return username, full_name, technician

@app.get("/admin/users")
def list_system_users(q: str | None = None, db: Session = Depends(get_db), _=Depends(require("ADMIN"))):
    query = select(User)
    if q:
        term = f"%{q.strip()}%"
        query = query.where((User.username.ilike(term)) | (User.full_name.ilike(term)))
    rows = db.scalars(query.order_by(User.role, User.full_name)).all()
    return {"total": len(rows), "items": [system_user_dict(row) for row in rows]}

@app.post("/admin/users", status_code=201)
def create_system_user(data: SystemUserIn, db: Session = Depends(get_db), u=Depends(require("ADMIN"))):
    if not data.password: raise HTTPException(422, "La contraseña inicial es obligatoria")
    username, full_name, technician = validate_system_user(data, db)
    user = User(username=username, full_name=full_name, password_hash=hash_pw(data.password),
                role=data.role, technician_id=technician.id if technician else None)
    if technician: technician.username = username
    db.add(user); db.flush()
    audit(db, u, "CREAR_USUARIO_SISTEMA", f"user:{user.id}", None,
          {"username": username, "full_name": full_name, "role": data.role,
           "technician_id": user.technician_id})
    db.commit(); db.refresh(user)
    return system_user_dict(user)

@app.put("/admin/users/{user_id}")
def update_system_user(user_id: int, data: SystemUserIn, db: Session = Depends(get_db), u=Depends(require("ADMIN"))):
    user = db.get(User, user_id)
    if not user: raise HTTPException(404, "Usuario no existe")
    if user.role == "ADMIN" and data.role != "ADMIN" and db.scalar(select(func.count(User.id)).where(User.role == "ADMIN")) <= 1:
        raise HTTPException(409, "No se puede quitar el último administrador del sistema")
    username, full_name, technician = validate_system_user(data, db, user_id)
    old = system_user_dict(user)
    old_technician = user.technician
    if old_technician and (not technician or technician.id != old_technician.id) and old_technician.username == user.username:
        old_technician.username = None
    user.username = username; user.full_name = full_name; user.role = data.role
    user.technician_id = technician.id if technician else None
    if technician: technician.username = username
    if data.password: user.password_hash = hash_pw(data.password)
    audit(db, u, "ACTUALIZAR_USUARIO_SISTEMA", f"user:{user.id}", old,
          {"username": username, "full_name": full_name, "role": data.role,
           "technician_id": user.technician_id, "password_reset": bool(data.password)})
    db.commit(); db.refresh(user)
    return system_user_dict(user)

class EmailSettingsIn(BaseModel):
    smtp_host: str = Field(default="", max_length=255)
    smtp_port: int = Field(default=587, ge=1, le=65535)
    smtp_user: str = Field(default="", max_length=255)
    smtp_password: str = Field(default="", max_length=1024)
    smtp_from: str = Field(default="", max_length=255)
    smtp_from_name: str = Field(default="Gestión Técnica", max_length=120)
    smtp_use_tls: bool = True
    notification_email_to: str = Field(default="", max_length=2000)
    report_email_to: str = Field(default="", max_length=2000)
    report_email_cc: str = Field(default="", max_length=2000)

class EmailTestIn(BaseModel):
    email: str

def normalized_addresses(value: str, label: str) -> str:
    addresses = [address.strip().lower() for address in re.split(r"[;,]", value) if address.strip()]
    invalid = [address for address in addresses if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", address)]
    if invalid: raise HTTPException(422, f"Revisa las direcciones de {label}: {', '.join(invalid)}")
    return ", ".join(dict.fromkeys(addresses))

@app.get("/admin/email-settings")
def get_email_settings(db: Session = Depends(get_db), _=Depends(require("ADMIN"))):
    return email_settings_dict(db)

@app.put("/admin/email-settings")
def update_email_settings(data: EmailSettingsIn, db: Session = Depends(get_db), u=Depends(require("ADMIN"))):
    fields = [data.smtp_host, data.smtp_user, data.smtp_from, data.smtp_from_name,
              data.notification_email_to, data.report_email_to, data.report_email_cc]
    if any("\r" in value or "\n" in value for value in fields):
        raise HTTPException(422, "No se permiten saltos de línea en la configuración de correo")
    sender = data.smtp_from.strip().lower()
    if sender and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", sender):
        raise HTTPException(422, "El correo remitente no es válido")
    old = email_settings_dict(db)
    settings = db.get(EmailSettings, 1)
    if not settings:
        settings = EmailSettings(id=1)
        db.add(settings)
    settings.smtp_host = data.smtp_host.strip()
    settings.smtp_port = data.smtp_port
    settings.smtp_user = data.smtp_user.strip()
    settings.smtp_from = sender
    settings.smtp_from_name = data.smtp_from_name.strip() or "Gestión Técnica"
    settings.smtp_use_tls = data.smtp_use_tls
    settings.notification_email_to = normalized_addresses(data.notification_email_to, "destinatarios de notificaciones")
    settings.report_email_to = normalized_addresses(data.report_email_to, "destinatarios de informes")
    settings.report_email_cc = normalized_addresses(data.report_email_cc, "copias de informes")
    settings.updated_by = u.id
    if data.smtp_password:
        settings.smtp_password_encrypted = email_cipher().encrypt(data.smtp_password.encode()).decode()
    new_values = {"smtp_host": settings.smtp_host, "smtp_port": settings.smtp_port,
                  "smtp_user": settings.smtp_user, "smtp_from": settings.smtp_from,
                  "smtp_from_name": settings.smtp_from_name, "smtp_use_tls": settings.smtp_use_tls,
                  "notification_email_to": settings.notification_email_to,
                  "report_email_to": settings.report_email_to, "report_email_cc": settings.report_email_cc,
                  "smtp_password_changed": bool(data.smtp_password)}
    audit(db, u, "ACTUALIZAR_CONFIGURACION_CORREO", "email_settings", old, new_values)
    db.commit()
    return email_settings_dict(db)

@app.post("/admin/email-settings/test")
def test_email_settings(payload: EmailTestIn, db: Session = Depends(get_db), u=Depends(require("ADMIN"))):
    recipient = normalized_addresses(payload.email, "correo de prueba")
    config = active_email_settings()
    if not config["smtp_host"] or not config["smtp_user"] or not config["smtp_password"]:
        raise HTTPException(503, "Completa el servidor SMTP, usuario y contraseña antes de enviar una prueba")
    message = EmailMessage()
    message["Subject"] = "Prueba de correo · Gestión Técnica"
    message["From"] = formataddr((config["smtp_from_name"], config["smtp_from"] or config["smtp_user"]))
    message["To"] = recipient
    message.set_content("La configuración de correo de Gestión Técnica funciona correctamente.")
    if not send_email_message(message, config):
        raise HTTPException(502, "No se pudo enviar la prueba. Revisa host, puerto, TLS y credenciales SMTP.")
    audit(db, u, "PROBAR_CONFIGURACION_CORREO", "email_settings", None, {"to": recipient, "sent": True})
    db.commit()
    return {"sent": True, "to": recipient}

def staff_user(u: User = Depends(current_user)) -> User:
    if u.role == "USUARIO": raise HTTPException(403, "Esta sección es solo para el equipo técnico")
    return u

TECHNICIAN_STATES = {"ACTIVO", "INACTIVO", "LICENCIA", "VACACIONES"}

class TechnicianIn(BaseModel):
    first_name: str
    last_name: str
    position: str | None = None
    username: str | None = None
    status: str = "ACTIVO"
    shift: str | None = None
    contracted_hours: int | None = None
    specialties: str | None = None
    hire_date: dt.date | None = None
    notes: str | None = None

@app.get("/technicians/{technician_id}/avatar")
def get_technician_avatar(technician_id: int, db: Session = Depends(get_db), _=Depends(staff_user)):
    technician = db.get(Technician, technician_id)
    if not technician: raise HTTPException(404, "Técnico no existe")
    if not technician.avatar_data: raise HTTPException(404, "El técnico no tiene una foto cargada")
    return Response(content=technician.avatar_data, media_type=technician.avatar_content_type,
                    headers={"Cache-Control": "private, no-store"})

@app.put("/technicians/{technician_id}/avatar")
async def update_technician_avatar(technician_id: int, file: UploadFile = File(...),
                                   db: Session = Depends(get_db), u=Depends(staff_user)):
    technician = db.get(Technician, technician_id)
    if not technician: raise HTTPException(404, "Técnico no existe")
    if u.role not in {"ADMIN", "JEFE"} and u.technician_id != technician_id:
        raise HTTPException(403, "Solo puedes cambiar tu propia foto")
    image, content_type = await read_avatar_image(file)
    technician.avatar_data = image
    technician.avatar_content_type = content_type
    audit(db, u, "ACTUALIZAR_AVATAR_TECNICO", f"technician:{technician_id}", None,
          {"content_type": content_type, "size": len(image)})
    db.commit()
    return {"has_avatar": True}

@app.get("/technicians")
def list_technicians(status: str | None = None, q: str | None = None,
                     page: int = Query(1, ge=1), size: int = Query(50, le=200),
                     db: Session = Depends(get_db), _=Depends(staff_user)):
    query = select(Technician)
    if status: query = query.where(Technician.status == status)
    if q:
        term = f"%{q.strip()}%"
        query = query.where((Technician.first_name.ilike(term)) | (Technician.last_name.ilike(term)) | (Technician.username.ilike(term)))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    rows = db.scalars(query.order_by(Technician.last_name, Technician.first_name)
                      .offset((page-1)*size).limit(size)).all()
    return {"total": total, "items": [technician_dict(t) for t in rows]}

def technician_dict(t):
    return {"id": t.id, "first_name": t.first_name, "last_name": t.last_name,
            "position": t.position, "username": t.username, "status": t.status,
            "shift": t.shift, "contracted_hours": t.contracted_hours,
            "specialties": t.specialties, "hire_date": t.hire_date, "notes": t.notes}

def validate_technician_assignment(db: Session, name: str | None, previous: str | None = None):
    normalized = (name or "").strip().casefold()
    if not normalized or normalized == (previous or "").strip().casefold():
        return
    matches = [tech for tech in db.scalars(select(Technician)).all()
               if f"{tech.first_name} {tech.last_name}".strip().casefold() == normalized]
    if matches and all(tech.status == "INACTIVO" for tech in matches):
        raise HTTPException(422, "El técnico está inactivo; selecciona otro técnico")

def validate_technician(t: TechnicianIn):
    if not t.first_name.strip() or not t.last_name.strip():
        raise HTTPException(422, "Nombre y apellido son obligatorios")
    if t.status not in TECHNICIAN_STATES:
        raise HTTPException(422, "Estado de técnico inválido")
    if t.contracted_hours is not None and t.contracted_hours < 0:
        raise HTTPException(422, "Las horas contratadas no pueden ser negativas")

@app.post("/technicians", status_code=201)
def create_technician(t: TechnicianIn, db: Session = Depends(get_db), u=Depends(require("JEFE", "ADMIN"))):
    validate_technician(t)
    username = t.username.strip() if t.username and t.username.strip() else None
    if username and db.scalar(select(Technician).where(Technician.username == username)):
        raise HTTPException(409, "El usuario ya está asociado a otro técnico")
    values = t.model_dump(); values["first_name"] = t.first_name.strip(); values["last_name"] = t.last_name.strip(); values["username"] = username
    tech = Technician(**values)
    db.add(tech); db.flush()
    audit(db, u, "CREAR_TECNICO", f"technician:{tech.id}", None, values)
    db.commit()
    return technician_dict(tech)

@app.put("/technicians/{technician_id}")
def update_technician(technician_id: int, t: TechnicianIn, db: Session = Depends(get_db), u=Depends(require("JEFE", "ADMIN"))):
    validate_technician(t)
    tech = db.get(Technician, technician_id)
    if not tech: raise HTTPException(404, "Técnico no existe")
    username = t.username.strip() if t.username and t.username.strip() else None
    if username and db.scalar(select(Technician).where(Technician.username == username, Technician.id != technician_id)):
        raise HTTPException(409, "El usuario ya está asociado a otro técnico")
    old = technician_dict(tech)
    values = t.model_dump(); values["first_name"] = t.first_name.strip(); values["last_name"] = t.last_name.strip(); values["username"] = username
    for key, value in values.items(): setattr(tech, key, value)
    audit(db, u, "ACTUALIZAR_TECNICO", f"technician:{tech.id}", old, values)
    db.commit()
    return technician_dict(tech)

class PartIn(BaseModel):
    code: str
    name: str
    category: str | None = None
    brand: str | None = None
    model: str | None = None
    initial_stock: int = 0
    minimum_stock: int = 0
    location: str | None = None
    unit_cost: float | None = None
    supplier: str | None = None
    inventory_type: str = "REPUESTO"

def part_dict(p: Part):
    return {"id": p.id, "code": p.code, "name": p.name, "category": p.category,
            "brand": p.brand, "model": p.model, "stock": p.stock,
            "minimum_stock": p.minimum_stock, "location": p.location,
            "unit_cost": float(p.unit_cost) if p.unit_cost is not None else None,
            "supplier": p.supplier, "status": p.status,
            "inventory_type": p.inventory_type,
            "low_stock": p.stock <= p.minimum_stock}

@app.get("/parts")
def list_parts(q: str | None = None, low_stock: bool = False, inventory_type: str | None = None,
               page: int = Query(1, ge=1), size: int = Query(50, le=200),
               db: Session = Depends(get_db), _=Depends(staff_user)):
    query = select(Part)
    if inventory_type:
        if inventory_type not in {"REPUESTO", "INSUMO"}:
            raise HTTPException(422, "Tipo de inventario inválido")
        query = query.where(Part.inventory_type == inventory_type)
    if q:
        term = f"%{q.strip()}%"
        query = query.where((Part.code.ilike(term)) | (Part.name.ilike(term)) | (Part.category.ilike(term)))
    if low_stock: query = query.where(Part.stock <= Part.minimum_stock)
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    rows = db.scalars(query.order_by(Part.name).offset((page-1)*size).limit(size)).all()
    return {"total": total, "items": [part_dict(p) for p in rows]}

@app.post("/parts", status_code=201)
def create_part(p: PartIn, db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    code = p.code.strip().upper()
    if not code or not p.name.strip(): raise HTTPException(422, "Código y nombre son obligatorios")
    if p.initial_stock < 0 or p.minimum_stock < 0: raise HTTPException(422, "El stock no puede ser negativo")
    if p.unit_cost is not None and p.unit_cost < 0: raise HTTPException(422, "El costo no puede ser negativo")
    if p.inventory_type not in {"REPUESTO", "INSUMO"}: raise HTTPException(422, "Tipo de inventario inválido")
    if db.scalar(select(Part).where(Part.code == code)): raise HTTPException(409, "El código ya existe")
    values = p.model_dump(exclude={"initial_stock"}); values["code"] = code; values["name"] = p.name.strip(); values["stock"] = p.initial_stock
    part = Part(**values); db.add(part); db.flush()
    if p.initial_stock:
        db.add(PartMovement(part_id=part.id, movement_type="ENTRADA", quantity=p.initial_stock,
                            stock_after=p.initial_stock, user_id=u.id, notes="Stock inicial"))
    audit(db, u, "CREAR_INSUMO" if part.inventory_type == "INSUMO" else "CREAR_REPUESTO",
          f"part:{part.id}", None, values)
    db.commit()
    return part_dict(part)

PART_CSV_HEADERS = {
    "codigo": "code", "nombre": "name", "categoria": "category", "marca": "brand",
    "modelo": "model", "stock_inicial": "initial_stock", "existencia_inicial": "initial_stock",
    "stock_minimo": "minimum_stock", "minimo": "minimum_stock", "ubicacion": "location",
    "costo_unitario": "unit_cost", "costo": "unit_cost", "proveedor": "supplier",
}
PART_CSV_LIMIT_BYTES = 5 * 1024 * 1024
PART_CSV_LIMIT_ROWS = 1000

def normalize_part_csv_header(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.strip().lower())
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[\s-]+", "_", value)

def parse_part_csv(data: bytes, inventory_type: str, db: Session):
    if len(data) > PART_CSV_LIMIT_BYTES:
        raise HTTPException(413, "El archivo CSV supera el límite de 5 MB")
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            content = data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "El archivo debe estar guardado como CSV UTF-8") from exc
    if not content.strip():
        raise HTTPException(422, "El archivo CSV está vacío")
    try:
        header_line = content.splitlines()[0]
        delimiter = max((";", ",", "\t"), key=header_line.count)
        if header_line.count(delimiter) == 0:
            raise HTTPException(422, "No se detectó un separador CSV válido")
        reader = csv.reader(io.StringIO(content), delimiter=delimiter)
        headers = next(reader)
        rows = list(reader)
    except (csv.Error, IndexError, StopIteration) as exc:
        raise HTTPException(422, "No se pudo leer el CSV. Descarga y usa la plantilla de ejemplo") from exc
    field_indexes = {}
    for index, header in enumerate(headers):
        field = PART_CSV_HEADERS.get(normalize_part_csv_header(header))
        if field and field not in field_indexes:
            field_indexes[field] = index
    missing = [label for field, label in (("code", "codigo"), ("name", "nombre"))
               if field not in field_indexes]
    if missing:
        raise HTTPException(422, f"Faltan columnas obligatorias: {', '.join(missing)}")

    records, errors, seen_codes = [], [], set()
    nonempty_count = 0
    for row_number, values in enumerate(rows, start=2):
        if not any((value or "").strip() for value in values):
            continue
        nonempty_count += 1
        if nonempty_count > PART_CSV_LIMIT_ROWS:
            raise HTTPException(413, f"El archivo supera el máximo de {PART_CSV_LIMIT_ROWS} filas")
        raw = {field: (values[index].strip() if index < len(values) else "")
               for field, index in field_indexes.items()}
        code = raw.get("code", "").strip().upper()
        name = raw.get("name", "").strip()
        problems = []
        if not code: problems.append("Falta el código")
        elif len(code) > 50: problems.append("El código supera 50 caracteres")
        if not name: problems.append("Falta el nombre")
        elif len(name) > 160: problems.append("El nombre supera 160 caracteres")
        values_out = {"code": code, "name": name, "inventory_type": inventory_type}
        limits = {"category": 80, "brand": 80, "model": 80, "location": 120, "supplier": 120}
        for field, limit in limits.items():
            value = raw.get(field, "").strip() or None
            if value and len(value) > limit: problems.append(f"{field}: máximo {limit} caracteres")
            values_out[field] = value
        for field, label in (("initial_stock", "stock inicial"), ("minimum_stock", "stock mínimo")):
            value = raw.get(field, "").strip() or "0"
            if not re.fullmatch(r"\d+", value):
                problems.append(f"{label}: usa un entero igual o mayor que cero")
                values_out[field] = 0
            else:
                values_out[field] = int(value)
        raw_cost = raw.get("unit_cost", "").strip()
        if raw_cost:
            cost_value = raw_cost.replace(" ", "")
            if "," in cost_value and "." in cost_value:
                cost_value = cost_value.replace(".", "").replace(",", ".")
            elif "," in cost_value:
                cost_value = cost_value.replace(",", ".")
            try:
                cost = float(cost_value)
                if not math.isfinite(cost) or cost < 0: raise ValueError
                values_out["unit_cost"] = cost
            except ValueError:
                problems.append("costo unitario: usa un número igual o mayor que cero")
                values_out["unit_cost"] = None
        else:
            values_out["unit_cost"] = None
        if code and code in seen_codes:
            problems.append("El código está repetido dentro del archivo")
        elif code:
            seen_codes.add(code)
        if problems:
            errors.append({"row": row_number, "code": code or None, "reason": "; ".join(problems)})
        else:
            records.append({"row": row_number, **values_out})

    existing = set()
    codes = [record["code"] for record in records]
    if codes:
        existing = {code.strip().upper() for code in db.scalars(select(Part.code).where(func.upper(func.trim(Part.code)).in_(codes))).all()}
    skipped = [{"row": record["row"], "code": record["code"]}
               for record in records if record["code"] in existing]
    valid = [record for record in records if record["code"] not in existing]
    return {"total_rows": nonempty_count, "valid": valid, "skipped": skipped, "errors": errors}

@app.post("/parts/import")
async def import_parts_csv(file: UploadFile = File(...), inventory_type: str = Form(...),
                          preview: bool = Form(True), db: Session = Depends(get_db),
                          u=Depends(require("JEFE"))):
    if inventory_type not in {"REPUESTO", "INSUMO"}:
        raise HTTPException(422, "Tipo de inventario inválido")
    if not file.filename or not file.filename.lower().endswith(".csv"):
        raise HTTPException(422, "Selecciona un archivo con extensión .csv")
    content = await file.read(PART_CSV_LIMIT_BYTES + 1)
    parsed = parse_part_csv(content, inventory_type, db)
    summary = {"total_rows": parsed["total_rows"], "valid_count": len(parsed["valid"]),
               "skipped_count": len(parsed["skipped"]), "skipped": parsed["skipped"][:30],
               "error_count": len(parsed["errors"]), "errors": parsed["errors"][:30],
               "sample": [{key: value for key, value in row.items() if key != "row"}
                          for row in parsed["valid"][:10]]}
    if preview:
        return {"preview": True, **summary}
    if not parsed["valid"]:
        raise HTTPException(422, "No hay filas nuevas válidas para cargar")
    created = []
    try:
        for values in parsed["valid"]:
            row = {key: value for key, value in values.items() if key != "row"}
            initial_stock = row.pop("initial_stock")
            part = Part(**row, stock=initial_stock)
            db.add(part)
            db.flush()
            if initial_stock:
                db.add(PartMovement(part_id=part.id, movement_type="ENTRADA", quantity=initial_stock,
                                    stock_after=initial_stock, user_id=u.id, notes="Stock inicial · carga CSV"))
            audit(db, u, "CARGA_CSV_INSUMO" if inventory_type == "INSUMO" else "CARGA_CSV_REPUESTO",
                  f"part:{part.id}", None, row | {"stock": initial_stock})
            created.append(part.code)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "Uno o más códigos se registraron mientras se procesaba el archivo. Vuelve a cargar para actualizar la vista previa") from exc
    return {"preview": False, "created_count": len(created), "created": created,
            "skipped_count": len(parsed["skipped"]), "skipped": parsed["skipped"][:30],
            "error_count": len(parsed["errors"]), "errors": parsed["errors"][:30]}

class PartMovementIn(BaseModel):
    movement_type: str
    quantity: int
    ticket_id: int | None = None
    machine: str | None = None
    technician: str | None = None
    notes: str | None = None

@app.post("/parts/{part_id}/movements", status_code=201)
def move_part(part_id: int, m: PartMovementIn, db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    validate_technician_assignment(db, m.technician)
    if m.movement_type not in {"ENTRADA", "SALIDA", "CONSUMO", "AJUSTE"}:
        raise HTTPException(422, "Tipo de movimiento inválido")
    if m.quantity == 0: raise HTTPException(422, "La cantidad no puede ser cero")
    if m.movement_type in {"ENTRADA", "SALIDA", "CONSUMO"} and m.quantity < 0:
        raise HTTPException(422, "La cantidad debe ser positiva")
    part = db.scalar(select(Part).where(Part.id == part_id).with_for_update())
    if not part: raise HTTPException(404, "Artículo de inventario no existe")
    if part.inventory_type == "INSUMO" and not (m.notes or "").strip():
        raise HTTPException(422, "Indica el motivo o destino del movimiento para mantener la trazabilidad de bodega")
    ticket = db.get(Ticket, m.ticket_id) if m.ticket_id else None
    if m.ticket_id and not ticket: raise HTTPException(422, "Ticket no existe")
    machine = db.scalar(select(Machine).where(Machine.number == m.machine)) if m.machine else None
    if m.machine and not machine: raise HTTPException(422, "Máquina no existe")
    if ticket and ticket.machine_id:
        if machine and machine.id != ticket.machine_id: raise HTTPException(422, "La máquina no coincide con el ticket")
        machine = machine or db.get(Machine, ticket.machine_id)
    if m.movement_type == "ENTRADA": delta = m.quantity
    elif m.movement_type in {"SALIDA", "CONSUMO"}: delta = -m.quantity
    else: delta = m.quantity  # Ajuste firmado: positivo suma, negativo resta.
    new_stock = part.stock + delta
    if new_stock < 0: raise HTTPException(409, "Stock insuficiente")
    old_stock = part.stock
    part.stock = new_stock
    movement = PartMovement(part_id=part.id, movement_type=m.movement_type, quantity=delta,
                            stock_after=new_stock, user_id=u.id, ticket_id=m.ticket_id,
                            machine_id=machine.id if machine else None, technician=m.technician, notes=m.notes)
    db.add(movement)
    audit(db, u, "MOVIMIENTO_INSUMO" if part.inventory_type == "INSUMO" else "MOVIMIENTO_REPUESTO", f"part:{part.id}", {"stock": old_stock},
          {"stock": new_stock, "movement_type": m.movement_type, "quantity": delta,
           "ticket_id": m.ticket_id, "machine": machine.number if machine else None})
    db.commit()
    if old_stock > part.minimum_stock >= new_stock:
        send_ticket_notification(f"Inventario bajo mínimo: {part.name}",
                                 f"El artículo {part.code} ({part.name}) quedó con {new_stock} unidades; mínimo configurado: {part.minimum_stock}.")
    return {"movement_id": movement.id, **part_dict(part)}

@app.get("/parts/{part_id}/movements")
def list_part_movements(part_id: int, page: int = Query(1, ge=1), size: int = Query(50, le=200),
                        db: Session = Depends(get_db), _=Depends(staff_user)):
    if not db.get(Part, part_id): raise HTTPException(404, "Artículo de inventario no existe")
    rows = db.scalars(select(PartMovement).where(PartMovement.part_id == part_id)
                      .order_by(PartMovement.at.desc()).offset((page-1)*size).limit(size)).all()
    return [{"id": r.id, "movement_type": r.movement_type, "quantity": r.quantity,
             "stock_after": r.stock_after, "at": r.at, "ticket_id": r.ticket_id,
             "machine": r.part_id and db.get(Machine, r.machine_id).number if r.machine_id else None,
             "technician": r.technician, "notes": r.notes} for r in rows]

class MachineIn(BaseModel):
    number: str
    island: str | None = None
    area: str | None = None
    manufacturer: str | None = None
    model: str | None = None
    serial: str | None = None

class MachinePositionIn(BaseModel):
    x: float
    y: float

@app.get("/floorplan")
def get_floorplan(db: Session = Depends(get_db), _=Depends(staff_user)):
    plan = db.get(FloorPlan, 1)
    if not plan: raise HTTPException(404, "No hay un plano de sala cargado")
    return {"filename": plan.filename, "updated_at": plan.updated_at,
            "size": len(plan.pdf_data)}

@app.get("/floorplan/image")
def get_floorplan_image(db: Session = Depends(get_db), _=Depends(staff_user)):
    plan = db.get(FloorPlan, 1)
    if not plan: raise HTTPException(404, "No hay un plano de sala cargado")
    return Response(content=plan.image_data, media_type="image/png",
                    headers={"Cache-Control": "no-cache, no-store, must-revalidate"})

@app.get("/floorplan/file")
def get_floorplan_file(db: Session = Depends(get_db), _=Depends(staff_user)):
    plan = db.get(FloorPlan, 1)
    if not plan: raise HTTPException(404, "No hay un plano de sala cargado")
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", plan.filename)
    return Response(content=plan.pdf_data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{safe_name}"; filename*=UTF-8\'\'{quote(plan.filename)}'})

@app.post("/floorplan")
async def upload_floorplan(file: UploadFile = File(...), db: Session = Depends(get_db),
                          u=Depends(require("JEFE"))):
    filename = Path((file.filename or "plano.pdf").replace("\\", "/")).name[:255]
    if not filename.lower().endswith(".pdf"):
        raise HTTPException(415, "Sube el plano en formato PDF")
    pdf_data = await file.read(25 * 1024 * 1024 + 1)
    if len(pdf_data) > 25 * 1024 * 1024:
        raise HTTPException(413, "El plano supera el límite de 25 MB")
    if not pdf_data.startswith(b"%PDF-"):
        raise HTTPException(415, "El archivo seleccionado no parece ser un PDF válido")
    try:
        image_data = render_floorplan(pdf_data)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    plan = db.get(FloorPlan, 1)
    old_filename = plan.filename if plan else None
    if plan:
        plan.filename = filename; plan.pdf_data = pdf_data; plan.image_data = image_data
        plan.updated_at = dt.datetime.now(dt.timezone.utc); plan.uploaded_by = u.id
    else:
        plan = FloorPlan(id=1, filename=filename, pdf_data=pdf_data,
                         image_data=image_data, uploaded_by=u.id)
        db.add(plan)
    audit(db, u, "ACTUALIZAR_PLANO_SALA", "floorplan", {"filename": old_filename},
          {"filename": filename, "size": len(pdf_data)})
    db.commit(); db.refresh(plan)
    return {"filename": plan.filename, "updated_at": plan.updated_at, "size": len(plan.pdf_data)}

@app.put("/machines/{number}/position")
def update_machine_position(number: str, position: MachinePositionIn,
                            db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    if not 0 <= position.x <= 1 or not 0 <= position.y <= 1:
        raise HTTPException(422, "La posición debe estar entre 0 y 1")
    machine = db.scalar(select(Machine).where(Machine.number == number))
    if not machine: raise HTTPException(404, "Máquina no existe")
    old = {"x": machine.position_x, "y": machine.position_y}
    machine.position_x = position.x; machine.position_y = position.y
    audit(db, u, "UBICAR_MAQUINA_EN_PLANO", f"machine:{number}", old,
          {"x": position.x, "y": position.y})
    db.commit()
    return {"number": machine.number, "x": machine.position_x, "y": machine.position_y}

@app.post("/machines", status_code=201)
def create_machine(m: MachineIn, db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    if db.scalar(select(Machine).where(Machine.number == m.number)):
        raise HTTPException(409, "Máquina ya existe")
    isl = None
    if m.island:
        isl = db.scalar(select(Island).where(Island.number == m.island)) or Island(number=m.island)
        db.add(isl)
    area = db.scalar(select(Area).where(Area.name == m.area)) if m.area else None
    mc = Machine(number=m.number, island=isl, area=area, manufacturer=m.manufacturer, model=m.model, serial=m.serial)
    db.add(mc); audit(db, u, "CREAR_MAQUINA", f"machine:{m.number}", None, m.model_dump())
    db.commit()
    return {"id": mc.id}

@app.put("/machines/{number}")
def update_machine(number: str, payload: MachineIn, db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    machine = db.scalar(select(Machine).where(Machine.number == number))
    if not machine:
        raise HTTPException(404, "Máquina no existe")
    values = {"island": payload.island.strip() if payload.island else None,
              "area": payload.area.strip() if payload.area else None,
              "manufacturer": payload.manufacturer.strip() if payload.manufacturer else None,
              "model": payload.model.strip() if payload.model else None,
              "serial": payload.serial.strip() if payload.serial else None}
    if values["island"] and len(values["island"]) > 10:
        raise HTTPException(422, "La isla no puede superar 10 caracteres")
    for field in ("manufacturer", "model", "serial"):
        if values[field] and len(values[field]) > 80:
            raise HTTPException(422, f"{field.capitalize()} no puede superar 80 caracteres")
    island = None
    if values["island"]:
        island = db.scalar(select(Island).where(Island.number == values["island"]))
        if not island:
            island = Island(number=values["island"]); db.add(island); db.flush()
    area = None
    if values["area"]:
        area = db.scalar(select(Area).where(Area.name == values["area"]))
        if not area:
            raise HTTPException(422, f"El área '{values['area']}' no existe en el catálogo")
    old = {"island": machine.island.number if machine.island else None,
           "area": machine.area.name if machine.area else None,
           "manufacturer": machine.manufacturer, "model": machine.model,
           "serial": machine.serial}
    machine.island = island; machine.area = area
    machine.manufacturer = values["manufacturer"]
    machine.model = values["model"]
    machine.serial = values["serial"]
    new = {"island": island.number if island else None, "area": area.name if area else None,
           "manufacturer": machine.manufacturer, "model": machine.model, "serial": machine.serial}
    audit(db, u, "ACTUALIZAR_MAQUINA", f"machine:{number}", old, new)
    db.commit()
    return {"number": machine.number, **new}

class MachineBulkDeleteIn(BaseModel):
    numbers: list[str]

@app.post("/machines/bulk-delete")
def bulk_delete_machines(payload: MachineBulkDeleteIn, db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    numbers = list(dict.fromkeys(number.strip() for number in payload.numbers if number.strip()))
    if not numbers:
        raise HTTPException(422, "Selecciona al menos una máquina")
    if len(numbers) > 200:
        raise HTTPException(422, "Puedes eliminar hasta 200 máquinas por operación")
    rows = db.scalars(select(Machine).where(Machine.number.in_(numbers))).all()
    by_number = {machine.number: machine for machine in rows}
    deleted = []
    skipped = []
    for number in numbers:
        machine = by_number.get(number)
        if not machine:
            continue
        references = []
        for label, model in (("tickets", Ticket), ("intervenciones", Intervention),
                             ("historial de estados", MachineStatusHistory), ("movimientos de repuestos", PartMovement),
                             ("planes preventivos", PreventiveMaintenancePlan),
                             ("historial preventivo", PreventiveMaintenanceLog)):
            if db.scalar(select(model.id).where(model.machine_id == machine.id).limit(1)) is not None:
                references.append(label)
        if references:
            skipped.append({"number": number, "references": references})
            continue
        audit(db, u, "ELIMINAR_MAQUINA", f"machine:{number}",
              {"number": machine.number, "manufacturer": machine.manufacturer,
               "model": machine.model, "serial": machine.serial, "status": machine.status}, None)
        db.delete(machine)
        deleted.append(number)
    db.commit()
    return {"deleted": deleted, "skipped": skipped,
            "not_found": [number for number in numbers if number not in by_number]}

def normalize_label(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "")
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", value.casefold())

def excel_text(value):
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    if isinstance(value, dt.datetime): return value.isoformat(sep=" ")
    return str(value).strip()

def read_machine_sheet(content: bytes, filename: str):
    if len(content) > 25 * 1024 * 1024: raise HTTPException(413, "El archivo supera el límite de 25 MB")
    suffix = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    try:
        if suffix == "xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            names = workbook.sheet_names()
            selected = next((n for n in names if "anexo3blistadodemaquinas" in normalize_label(n)), None)
            if selected is None: raise HTTPException(422, "No se encontró la hoja Anexo 3B- Listado de maquinas")
            sheet = workbook.sheet_by_name(selected)
            rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        elif suffix in {"xlsx", "xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            selected = next((n for n in workbook.sheetnames if "anexo3blistadodemaquinas" in normalize_label(n)), None)
            if selected is None: raise HTTPException(422, "No se encontró la hoja Anexo 3B- Listado de maquinas")
            sheet = workbook[selected]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        else:
            raise HTTPException(415, "Formato no compatible. Sube un archivo .xls, .xlsx o .xlsm")
    except HTTPException: raise
    except Exception as exc:
        raise HTTPException(422, f"No fue posible leer el archivo: {exc}")
    # Fila 10 es la referencia indicada. Si allí hay solo títulos, detecta encabezados cercanos.
    header_idx = 9
    for candidate in range(9, min(len(rows), 15)):
        labels = [normalize_label(excel_text(v)) for v in rows[candidate]]
        recognizable = sum(any(word in label for word in ("maquina", "equipo", "numero", "isla", "area", "sector", "fabricante", "modelo", "serie")) for label in labels)
        if recognizable >= 2:
            header_idx = candidate; break
    if header_idx >= len(rows): raise HTTPException(422, "El archivo no contiene datos desde la fila 10")
    headers = rows[header_idx]
    columns = [{"index": i, "label": excel_text(v) or f"Columna {i+1}"} for i, v in enumerate(headers)]
    populated = []
    for idx, row in enumerate(rows[header_idx+1:], start=header_idx+2):
        if any(excel_text(v) for v in row):
            populated.append({"row_number": idx, "cells": [excel_text(v) for v in row]})
    return selected, header_idx + 1, columns, populated

@app.post("/machines/import/preview")
async def preview_machine_import(file: UploadFile = File(...), _=Depends(staff_user)):
    sheet, header_row, columns, rows = read_machine_sheet(await file.read(), file.filename or "")
    suggested = {}
    aliases = {"number": ("numeromaquina", "maquina", "numerodeequipo", "numeroequipo", "equipo"),
               "island": ("isla", "numerodeisla"), "area": ("area", "sector", "nombresector", "sectordeubicacion"), "manufacturer": ("fabricante", "marca"),
               "model": ("modelo",), "serial": ("numerodeserie", "serie", "serial")}
    for field, options in aliases.items():
        col = next((c for c in columns if normalize_label(c["label"]) in options), None)
        if col: suggested[field] = col["index"]
    return {"sheet": sheet, "header_row": header_row, "data_start_row": header_row+1,
            "columns": columns, "suggested_mapping": suggested, "total_rows": len(rows), "preview": rows[:12]}

@app.post("/machines/import")
async def import_machines(file: UploadFile = File(...), column_map: str = Form(...),
                          db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    try: mapping = json.loads(column_map)
    except Exception: raise HTTPException(422, "Mapeo de columnas inválido")
    if not isinstance(mapping, dict) or "number" not in mapping:
        raise HTTPException(422, "Debes seleccionar la columna del número de máquina")
    sheet, header_row, columns, rows = read_machine_sheet(await file.read(), file.filename or "")
    valid_indexes = {c["index"] for c in columns}
    for field, index in mapping.items():
        if field not in {"number", "island", "area", "manufacturer", "model", "serial"} or index is not None and index not in valid_indexes:
            raise HTTPException(422, "El mapeo contiene una columna inválida")
    imported, errors, seen = [], [], set()
    created_count = 0
    updated_count = 0
    unchanged_count = 0
    area_rows = db.scalars(select(Area)).all()
    area_by_normalized = {normalize_label(area.name): area for area in area_rows}
    for row in rows:
        cells = row["cells"]
        original = {field: cells[index] if index is not None and index < len(cells) else "" for field, index in mapping.items()}
        number = original.get("number", "").strip()
        if not number:
            errors.append({"row": row["row_number"], "reason": "Falta número de máquina"}); continue
        if len(number) > 20:
            errors.append({"row": row["row_number"], "number": number, "reason": "Número supera 20 caracteres"}); continue
        if number in seen:
            errors.append({"row": row["row_number"], "number": number, "reason": "Número de máquina repetido dentro del archivo"}); continue
        seen.add(number)
        area_value = original.get("area", "").strip()
        area = area_by_normalized.get(normalize_label(area_value)) if area_value else None
        island_value = original.get("island", "").strip()
        warnings = []
        if island_value and len(island_value) > 10:
            errors.append({"row": row["row_number"], "number": number, "reason": "Isla supera 10 caracteres"}); continue
        limits = {"manufacturer": 80, "model": 80, "serial": 80}
        too_long = next(((field, original.get(field, "").strip()) for field, limit in limits.items()
                         if len(original.get(field, "").strip()) > limit), None)
        if too_long:
            errors.append({"row": row["row_number"], "number": number,
                           "reason": f"{too_long[0]} supera {limits[too_long[0]]} caracteres"}); continue
        if area_value and len(area_value) > 60:
            errors.append({"row": row["row_number"], "number": number, "reason": "Área / sector supera 60 caracteres"}); continue
        if area_value and not area:
            area = Area(name=area_value)
            db.add(area)
            db.flush()
            area_by_normalized[normalize_label(area_value)] = area
        machine = db.scalar(select(Machine).where(Machine.number == number))
        is_new = machine is None
        if machine is None:
            machine = Machine(number=number)
            db.add(machine)
            db.flush()
        old_values = {"island": machine.island.number if machine.island else None,
                      "area": machine.area.name if machine.area else None,
                      "manufacturer": machine.manufacturer, "model": machine.model,
                      "serial": machine.serial}
        if island_value:
            island = db.scalar(select(Island).where(Island.number == island_value))
            if not island:
                island = Island(number=island_value); db.add(island); db.flush()
            machine.island = island
        if area_value and area:
            machine.area = area
        for field in ("manufacturer", "model", "serial"):
            value = original.get(field, "").strip()
            if value:
                setattr(machine, field, value)
        new_values = {"island": machine.island.number if machine.island else None,
                      "area": machine.area.name if machine.area else None,
                      "manufacturer": machine.manufacturer, "model": machine.model,
                      "serial": machine.serial}
        action = "created" if is_new else "updated" if new_values != old_values else "unchanged"
        if action == "created": created_count += 1
        elif action == "updated": updated_count += 1
        else: unchanged_count += 1
        if action != "unchanged":
            audit(db, u, "IMPORTAR_MAQUINA" if is_new else "ACTUALIZAR_MAQUINA_IMPORTADA",
                  f"machine:{number}", None if is_new else old_values,
                  {**new_values, "source_file": file.filename,
                   "source_sheet": sheet, "source_row": row["row_number"], "action": action})
        imported.append({"row": row["row_number"], "number": number, "action": action,
                         "warning": "; ".join(warnings) or None})
    db.commit()
    return {"sheet": sheet, "header_row": header_row, "processed": len(rows), "imported": len(imported),
            "created": created_count, "updated": updated_count, "unchanged": unchanged_count,
            "errors_count": len(errors), "items": imported, "errors": errors}

@app.get("/machines")
def list_machines(status: str | None = None, page: int = Query(1, ge=1), size: int = Query(50, le=200),
                  db: Session = Depends(get_db), _=Depends(staff_user)):
    q = select(Machine)
    if status: q = q.where(Machine.status == status)
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    rows = db.scalars(q.order_by(Machine.number).offset((page-1)*size).limit(size)).all()
    return {"total": total, "items": [{"number": r.number, "status": r.status,
            "island": r.island.number if r.island else None, "area": r.area.name if r.area else None,
            "manufacturer": r.manufacturer, "model": r.model, "serial": r.serial,
            "position_x": r.position_x, "position_y": r.position_y} for r in rows]}

class IslandIn(BaseModel):
    number: str
    hall: str | None = None

@app.get("/islands")
def list_islands(q: str | None = None, page: int = Query(1, ge=1), size: int = Query(100, le=200),
                 db: Session = Depends(get_db), _=Depends(staff_user)):
    query = select(Island)
    if q:
        term = f"%{q.strip()}%"
        query = query.where((Island.number.ilike(term)) | (Island.hall.ilike(term)))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    islands = db.scalars(query.order_by(Island.hall, Island.number)
                         .offset((page-1)*size).limit(size)).all()
    items = []
    for island in islands:
        machines = db.scalars(select(Machine).where(Machine.island_id == island.id).order_by(Machine.number)).all()
        items.append({"id": island.id, "number": island.number, "hall": island.hall,
                      "machines": [{"number": machine.number, "status": machine.status,
                                    "area": machine.area.name if machine.area else None}
                                   for machine in machines]})
    return {"total": total, "items": items}

@app.post("/islands", status_code=201)
def create_island(i: IslandIn, db: Session = Depends(get_db), u=Depends(require("JEFE"))):
    number = i.number.strip()
    if not number: raise HTTPException(422, "El número de isla es obligatorio")
    if db.scalar(select(Island).where(Island.number == number)):
        raise HTTPException(409, "La isla ya existe")
    island = Island(number=number, hall=i.hall.strip() if i.hall and i.hall.strip() else None)
    db.add(island); db.flush()
    audit(db, u, "CREAR_ISLA", f"island:{island.id}", None, {"number": number, "hall": island.hall})
    db.commit()
    return {"id": island.id, "number": island.number, "hall": island.hall, "machines": []}

class StatusChange(BaseModel):
    new_status: str
    reason: str
    ticket_id: int | None = None

@app.post("/machines/{number}/status")
def change_status(number: str, c: StatusChange, db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    if c.new_status not in MACHINE_STATES: raise HTTPException(422, "Estado inválido")
    if not c.reason.strip(): raise HTTPException(422, "El motivo es obligatorio")
    m = db.scalar(select(Machine).where(Machine.number == number))
    if not m: raise HTTPException(404, "Máquina no existe")
    if m.status == c.new_status: raise HTTPException(409, "La máquina ya está en ese estado")
    now_ = dt.datetime.now(dt.timezone.utc)
    downtime = None
    if c.new_status == "OPERATIVA":  # cierra downtime abierto
        op = db.scalar(select(MachineStatusHistory).where(
            MachineStatusHistory.machine_id == m.id, MachineStatusHistory.new_status == "FUERA_DE_SERVICIO",
            MachineStatusHistory.downtime_end.is_(None)).order_by(MachineStatusHistory.at.desc()))
        if op:
            op.downtime_end = now_
            op.downtime_minutes = int((now_ - op.at).total_seconds() // 60)
            downtime = op.downtime_minutes
    db.add(MachineStatusHistory(machine_id=m.id, old_status=m.status, new_status=c.new_status,
                                at=now_, user_id=u.id, reason=c.reason, ticket_id=c.ticket_id))
    audit(db, u, "CAMBIO_ESTADO", f"machine:{number}", m.status, c.new_status)
    m.status = c.new_status
    db.commit()
    return {"number": number, "status": m.status, "downtime_minutes": downtime}

@app.get("/machines/{number}/history")
def history(number: str, db: Session = Depends(get_db), _=Depends(staff_user)):
    m = db.scalar(select(Machine).where(Machine.number == number))
    if not m: raise HTTPException(404, "Máquina no existe")
    status_rows = db.scalars(select(MachineStatusHistory).where(MachineStatusHistory.machine_id == m.id)
                             .order_by(MachineStatusHistory.at.desc())).all()
    movement_rows = db.scalars(select(PartMovement).where(PartMovement.machine_id == m.id)
                               .order_by(PartMovement.at.desc())).all()
    intervention_rows = db.scalars(select(Intervention).options(selectinload(Intervention.follow_ups).selectinload(InterventionFollowUp.user)).where(or_(
                                   Intervention.machine_id == m.id,
                                   (Intervention.machine_id.is_(None)) &
                                   Intervention.ticket_id.in_(select(Ticket.id).where(Ticket.machine_id == m.id))))
                                   .order_by(Intervention.occurred_at.desc())).all()
    history_rows = [{"type": "estado", "at": r.at, "old": r.old_status, "new": r.new_status,
                     "reason": r.reason, "ticket_id": r.ticket_id,
                     "downtime_minutes": r.downtime_minutes} for r in status_rows]
    history_rows.extend({"type": "inventario", "at": r.at, "inventory_type": r.part.inventory_type,
                         "part_name": r.part.name, "part_code": r.part.code,
                         "movement_type": r.movement_type, "quantity": r.quantity,
                         "stock_after": r.stock_after, "technician": r.technician,
                         "notes": r.notes, "ticket_id": r.ticket_id}
                        for r in movement_rows)
    history_rows.extend({"type": "intervencion", "at": r.occurred_at,
                         "work_type": r.work_type, "task": r.task, "detail": r.detail,
                         "technician": r.technician, "shift": r.shift,
                         "pending": r.follow_ups[-1].status != "RESUELTA" if r.follow_ups else r.pending,
                         "result": r.result, "notes": r.notes, "ticket_id": r.ticket_id,
                         "jira_number": r.jira_number, "area": r.area.name if r.area else None}
                        for r in intervention_rows)
    history_rows.sort(key=lambda row: row["at"], reverse=True)
    return history_rows

class TicketIn(BaseModel):
    task: str
    detail: str = ""
    machine: str | None = None
    technician: str | None = None
    shift: str | None = None
    priority: str = "NORMAL"
    status: str = "NUEVO"
    jira_number: str | None = None

TICKET_STATES = {"NUEVO", "ASIGNADO", "EN PROCESO", "PENDIENTE", "ESPERANDO REPUESTO", "ESPERANDO PROVEEDOR", "RESUELTO", "CERRADO", "CANCELADO"}
TICKET_PRIORITIES = {"BAJA", "NORMAL", "ALTA", "CRÍTICA"}

def ticket_dict(t: Ticket):
    sla_terminal = t.status in {"RESUELTO", "CERRADO", "CANCELADO"}
    if not t.sla_due_at:
        sla_status = "SIN SLA HISTÓRICO"
    elif sla_terminal:
        sla_status = "FINALIZADO"
    else:
        sla_status = "VENCIDO" if dt.datetime.now(dt.timezone.utc) >= t.sla_due_at else "EN PLAZO"
    return {"id": t.id, "created_at": t.created_at, "task": t.task, "detail": t.detail,
            "machine": t.machine.number if t.machine else None,
            "area": t.area.name if t.area else None, "status": t.status,
            "priority": t.priority, "jira": t.jira_number, "technician": t.technician,
            "sla_due_at": t.sla_due_at, "sla_status": sla_status,
            "shift": t.shift, "result": t.result,
            "requester": {"name": t.requester.full_name, "email": t.requester.username} if t.requester else None,
            "attachments": [{"id": attachment.id, "filename": attachment.filename,
                             "content_type": attachment.content_type}
                            for attachment in t.attachments]}

@app.post("/tickets", status_code=201)
async def create_ticket(task: str = Form(...), detail: str = Form(""), machine: str | None = Form(None),
                        images: list[UploadFile] = File(default=[]), db: Session = Depends(get_db),
                        u=Depends(current_user)):
    if u.role != "USUARIO":
        raise HTTPException(403, "Los tickets deben crearse desde el portal de usuarios")
    if not task.strip(): raise HTTPException(422, "La tarea es obligatoria")
    if len(images) > 10: raise HTTPException(422, "Puedes adjuntar hasta 10 imágenes por ticket")
    image_payloads = []
    allowed_signatures = {"image/jpeg": (b"\xff\xd8\xff",), "image/png": (b"\x89PNG\r\n\x1a\n",),
                          "image/gif": (b"GIF87a", b"GIF89a"), "image/webp": (b"RIFF",)}
    for image in images:
        if not image.filename:
            continue
        content = await image.read(2 * 1024 * 1024 + 1)
        if len(content) > 2 * 1024 * 1024:
            raise HTTPException(413, f"La imagen {Path(image.filename or 'archivo').name} supera 2 MB")
        content_type = (image.content_type or "").lower()
        signatures = allowed_signatures.get(content_type)
        valid_signature = bool(signatures and any(content.startswith(signature) for signature in signatures))
        if content_type == "image/webp": valid_signature = content.startswith(b"RIFF") and content[8:12] == b"WEBP"
        if not valid_signature:
            raise HTTPException(415, f"{Path(image.filename or 'archivo').name} no es una imagen compatible (JPG, PNG, GIF o WebP)")
        image_payloads.append((Path(image.filename or "imagen").name[:255], content_type, content))
    status = "NUEVO"
    priority = "NORMAL"
    if status not in TICKET_STATES: raise HTTPException(422, "Estado de ticket inválido")
    if priority not in TICKET_PRIORITIES: raise HTTPException(422, "Prioridad inválida")
    mid = None
    machine_value = (machine or "").strip()
    if machine_value:
        m = db.scalar(select(Machine).where(Machine.number == machine_value))
        if not m: raise HTTPException(422, "Máquina no existe")
        machine = m
        mid = m.id
    created_at = dt.datetime.now(dt.timezone.utc)
    tk = Ticket(machine_id=mid, area_id=machine.area_id if machine else None,
                created_at=created_at, sla_due_at=created_at + dt.timedelta(hours=TICKET_SLA_HOURS[priority]),
                created_by=u.id, task=task.strip(), detail=detail.strip(), technician=None,
                shift=None, status=status, priority=priority, jira_number=None)
    db.add(tk); db.flush()
    for filename, content_type, content in image_payloads:
        db.add(TicketAttachment(ticket_id=tk.id, filename=filename, content_type=content_type,
                                image_data=content))
    db.add(TicketStatusHistory(ticket_id=tk.id, old_status="", new_status=tk.status, user_id=u.id, note="Ticket creado"))
    audit(db, u, "CREAR_TICKET", f"ticket:{tk.id}", None,
          {"task": task.strip(), "detail": detail.strip(), "machine": machine_value or None,
           "attachments": len(image_payloads)})
    db.commit()
    try:
        send_ticket_notification(f"Nuevo ticket #{tk.id}: {tk.task}",
                                 f"Se registró una nueva solicitud.\nSolicitante: {u.full_name or u.username}\nTarea: {tk.task}\nMáquina: {machine_value or 'No indicada'}\n\nIngresa a Gestión Técnica para revisarla.")
    except Exception:
        logger.exception("El ticket #%s se guardó, pero falló la notificación por correo", tk.id)
    return {"id": tk.id}

@app.get("/tickets/{ticket_id}/attachments/{attachment_id}")
def get_ticket_attachment(ticket_id: int, attachment_id: int, db: Session = Depends(get_db),
                          u=Depends(current_user)):
    ticket = db.get(Ticket, ticket_id)
    if not ticket or u.role == "USUARIO" and ticket.created_by != u.id:
        raise HTTPException(404, "Ticket no existe")
    attachment = db.scalar(select(TicketAttachment).where(
        TicketAttachment.id == attachment_id, TicketAttachment.ticket_id == ticket_id))
    if not attachment: raise HTTPException(404, "Imagen no existe")
    return Response(content=attachment.image_data, media_type=attachment.content_type,
                    headers={"Content-Disposition": f'inline; filename="{attachment.filename}"',
                             "X-Content-Type-Options": "nosniff"})

@app.get("/tickets")
def list_tickets(status: str | None = None, priority: str | None = None, search_text: str | None = None,
                 attention: str | None = None,
                 page: int = Query(1, ge=1), size: int = Query(25, le=100),
                 db: Session = Depends(get_db), u=Depends(current_user)):
    query = select(Ticket)
    if u.role == "USUARIO": query = query.where(Ticket.created_by == u.id)
    if status: query = query.where(Ticket.status == status)
    if priority: query = query.where(Ticket.priority == priority)
    open_tickets = Ticket.status.notin_({"RESUELTO", "CERRADO", "CANCELADO"})
    if attention == "overdue":
        query = query.where(open_tickets, Ticket.sla_due_at.is_not(None),
                            Ticket.sla_due_at <= dt.datetime.now(dt.timezone.utc))
    elif attention == "unassigned":
        query = query.where(open_tickets, (Ticket.technician.is_(None)) | (func.trim(Ticket.technician) == ""))
    elif attention == "urgent":
        query = query.where(open_tickets, Ticket.priority.in_({"CRÍTICA", "ALTA"}))
    elif attention == "under_24h":
        query = query.where(open_tickets, Ticket.created_at > dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=24))
    elif attention == "24_to_72h":
        attention_now = dt.datetime.now(dt.timezone.utc)
        query = query.where(open_tickets, Ticket.created_at <= attention_now - dt.timedelta(hours=24),
                            Ticket.created_at > attention_now - dt.timedelta(hours=72))
    elif attention == "over_72h":
        query = query.where(open_tickets, Ticket.created_at <= dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=72))
    elif attention == "over_24h":
        query = query.where(open_tickets, Ticket.created_at <= dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=24))
    if search_text:
        term = f"%{search_text.strip()}%"
        query = query.where((Ticket.task.ilike(term)) | (Ticket.detail.ilike(term)) | (Ticket.technician.ilike(term)) | (Ticket.jira_number.ilike(term)))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    rows = db.scalars(query.order_by(Ticket.created_at.desc()).offset((page-1)*size).limit(size)).all()
    return {"total": total, "items": [ticket_dict(r) for r in rows]}

@app.get("/tickets/{ticket_id}")
def get_ticket(ticket_id: int, db: Session = Depends(get_db), u=Depends(current_user)):
    ticket = db.get(Ticket, ticket_id)
    if not ticket or u.role == "USUARIO" and ticket.created_by != u.id:
        raise HTTPException(404, "Ticket no existe")
    return ticket_dict(ticket)

@app.get("/tickets/{ticket_id}/inventory")
def ticket_inventory_usage(ticket_id: int, db: Session = Depends(get_db), _=Depends(staff_user)):
    if not db.get(Ticket, ticket_id): raise HTTPException(404, "Ticket no existe")
    rows = db.scalars(select(PartMovement).options(selectinload(PartMovement.part)).where(
        PartMovement.ticket_id == ticket_id, PartMovement.movement_type == "CONSUMO")
        .order_by(PartMovement.at.desc(), PartMovement.id.desc()).limit(100)).all()
    return [{"id": row.id, "at": row.at, "code": row.part.code, "name": row.part.name,
             "inventory_type": row.part.inventory_type, "quantity": abs(row.quantity),
             "stock_after": row.stock_after, "technician": row.technician, "notes": row.notes}
            for row in rows]

class InventoryUseIn(BaseModel):
    part_id: int = Field(gt=0)
    quantity: int = Field(gt=0, le=1000000)

class TicketUpdate(BaseModel):
    status: str | None = None
    priority: str | None = None
    technician: str | None = None
    shift: str | None = None
    result: str | None = None
    note: str | None = None
    parts_used: list[InventoryUseIn] = Field(default_factory=list, max_length=100)

def lock_inventory_for_usage(lines: list[InventoryUseIn], db: Session):
    part_ids = [line.part_id for line in lines]
    if len(part_ids) != len(set(part_ids)):
        raise HTTPException(422, "No repitas el mismo artículo; suma las cantidades en una sola línea")
    if not part_ids:
        return {}
    parts = {part.id: part for part in db.scalars(
        select(Part).where(Part.id.in_(part_ids)).order_by(Part.id).with_for_update()).all()}
    missing = [str(part_id) for part_id in part_ids if part_id not in parts]
    if missing:
        raise HTTPException(422, f"Artículo(s) de inventario no existe(n): {', '.join(missing)}")
    for line in lines:
        part = parts[line.part_id]
        if line.quantity > part.stock:
            raise HTTPException(409, f"Stock insuficiente para {part.code} · {part.name}: disponibles {part.stock}, solicitados {line.quantity}")
    return parts

def record_inventory_usage(lines: list[InventoryUseIn], parts: dict, db: Session, user: User,
                           *, ticket_id: int | None, machine_id: int | None, technician: str | None,
                           reference: str, audit_action: str):
    low_stock_notifications = []
    for line in lines:
        part = parts[line.part_id]
        old_stock = part.stock
        part.stock -= line.quantity
        db.add(PartMovement(part_id=part.id, movement_type="CONSUMO", quantity=-line.quantity,
                            stock_after=part.stock, user_id=user.id, ticket_id=ticket_id,
                            machine_id=machine_id, technician=technician or user.full_name or user.username,
                            notes=reference[:4000]))
        audit(db, user, audit_action, f"part:{part.id}", {"stock": old_stock},
              {"stock": part.stock, "quantity": line.quantity, "ticket_id": ticket_id,
               "machine_id": machine_id, "inventory_type": part.inventory_type})
        if old_stock > part.minimum_stock >= part.stock:
            low_stock_notifications.append((part.code, part.name, part.stock, part.minimum_stock))
    return low_stock_notifications

def notify_inventory_below_minimum(notifications: list[tuple[str, str, int, int]]):
    for code, name, stock, minimum in notifications:
        try:
            send_ticket_notification(f"Inventario bajo mínimo: {name}",
                                     f"El artículo {code} ({name}) quedó con {stock} unidades; mínimo configurado: {minimum}.")
        except Exception:
            logger.exception("No se pudo notificar stock bajo mínimo para %s", code)

@app.patch("/tickets/{ticket_id}")
def update_ticket(ticket_id: int, update: TicketUpdate, db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    ticket = db.get(Ticket, ticket_id)
    if not ticket: raise HTTPException(404, "Ticket no existe")
    values = update.model_dump(exclude_unset=True)
    if "technician" in values:
        validate_technician_assignment(db, values["technician"], ticket.technician)
    note = values.pop("note", None)
    parts_used = values.pop("parts_used", [])
    inventory_parts = lock_inventory_for_usage(parts_used, db)
    if "status" in values and values["status"] not in TICKET_STATES: raise HTTPException(422, "Estado de ticket inválido")
    if "priority" in values and values["priority"] not in TICKET_PRIORITIES: raise HTTPException(422, "Prioridad inválida")
    if values.get("status") in {"RESUELTO", "CERRADO"} and values["status"] != ticket.status:
        resolution = values.get("result", ticket.result)
        if not resolution or not resolution.strip():
            raise HTTPException(422, "Registra la solución o causa antes de resolver o cerrar el ticket")
    old = {key: getattr(ticket, key) for key in values}
    old_technician = ticket.technician
    if "priority" in values and values["priority"] != ticket.priority and ticket.sla_due_at:
        old["sla_due_at"] = ticket.sla_due_at
        values["sla_due_at"] = ticket.created_at + dt.timedelta(hours=TICKET_SLA_HOURS[values["priority"]])
        ticket.sla_approaching_notified_at = None
        ticket.sla_overdue_notified_at = None
    if values.get("status") and values["status"] != ticket.status:
        db.add(TicketStatusHistory(ticket_id=ticket.id, old_status=ticket.status, new_status=values["status"], user_id=u.id, note=note))
    for key, value in values.items():
        if isinstance(value, str): value = value.strip() or None
        setattr(ticket, key, value)
    low_stock_notifications = record_inventory_usage(
        parts_used, inventory_parts, db, u, ticket_id=ticket.id,
        machine_id=ticket.machine_id, technician=ticket.technician,
        reference=f"Consumo para Ticket #{ticket.id} · {ticket.task}",
        audit_action="CONSUMO_INVENTARIO_DESDE_TICKET")
    audit(db, u, "ACTUALIZAR_TICKET", f"ticket:{ticket.id}", old, values)
    db.commit()
    notify_inventory_below_minimum(low_stock_notifications)
    if ticket.technician and ticket.technician != old_technician:
        technician = db.scalar(select(Technician).where(
            func.lower(func.trim(Technician.first_name + " " + Technician.last_name)) == ticket.technician.lower().strip()))
        account = db.scalar(select(User).where(User.technician_id == technician.id)) if technician else None
        recipient = account.username if account and re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", account.username) else None
        if recipient:
            send_system_email([recipient], f"Ticket asignado · #{ticket.id}",
                f"Se te asignó el ticket #{ticket.id}.\nTarea: {ticket.task}\nPrioridad: {ticket.priority}\nVencimiento SLA: {ticket.sla_due_at:%d/%m/%Y %H:%M UTC}" if ticket.sla_due_at else f"Se te asignó el ticket #{ticket.id}.\nTarea: {ticket.task}\nPrioridad: {ticket.priority}")
    if "status" in values and values["status"] != old.get("status"):
        requester = db.get(User, ticket.created_by)
        send_system_email([requester.username] if requester else [], f"Actualización del ticket #{ticket.id}",
                          f"El estado de tu solicitud '{ticket.task}' cambió a {ticket.status}.\n{note or ''}")
    return ticket_dict(ticket)

class MaintenancePlanIn(BaseModel):
    machine_number: str | None = Field(default=None, min_length=1, max_length=20)
    machine_numbers: list[str] = Field(default_factory=list)
    title: str = Field(min_length=1, max_length=160)
    description: str = Field(default="", max_length=4000)
    interval_days: int = Field(ge=1, le=3650)
    next_due_at: dt.date

class MaintenanceCompleteIn(BaseModel):
    technician: str = Field(default="", max_length=120)
    notes: str = Field(default="", max_length=4000)

def maintenance_plan_dict(plan: PreventiveMaintenancePlan):
    return {"id": plan.id, "machine_number": plan.machine.number, "machine_id": plan.machine_id,
            "island": plan.machine.island.number if plan.machine.island else None,
            "title": plan.title, "description": plan.description, "interval_days": plan.interval_days,
            "next_due_at": plan.next_due_at, "active": plan.active, "created_at": plan.created_at}

@app.get("/maintenance/plans")
def list_maintenance_plans(db: Session = Depends(get_db), _=Depends(staff_user)):
    plans = db.scalars(select(PreventiveMaintenancePlan).order_by(
        PreventiveMaintenancePlan.active.desc(), PreventiveMaintenancePlan.next_due_at)).all()
    return [maintenance_plan_dict(plan) for plan in plans]

@app.post("/maintenance/plans", status_code=201)
def create_maintenance_plan(payload: MaintenancePlanIn, db: Session = Depends(get_db), u=Depends(require("JEFE", "SUPERVISOR"))):
    numbers = list(dict.fromkeys(number.strip() for number in [*payload.machine_numbers, payload.machine_number or ""] if number.strip()))
    if not numbers: raise HTTPException(422, "Selecciona al menos una máquina")
    found = db.scalars(select(Machine).where(Machine.number.in_(numbers))).all()
    machines_by_number = {machine.number: machine for machine in found}
    missing = [number for number in numbers if number not in machines_by_number]
    if missing: raise HTTPException(404, f"Máquina(s) no existe(n): {', '.join(missing)}")
    due = dt.datetime.combine(payload.next_due_at, dt.time(hour=12), tzinfo=dt.timezone.utc)
    plans = []
    for number in numbers:
        machine = machines_by_number[number]
        plan = PreventiveMaintenancePlan(machine_id=machine.id, title=payload.title.strip(),
            description=payload.description.strip(), interval_days=payload.interval_days, next_due_at=due, created_by=u.id)
        db.add(plan); db.flush()
        audit(db, u, "CREAR_PLAN_MANTENIMIENTO", f"maintenance:{plan.id}", None, maintenance_plan_dict(plan))
        plans.append(plan)
    db.commit(); db.refresh(plan)
    return [maintenance_plan_dict(item) for item in plans]

@app.post("/maintenance/plans/{plan_id}/complete")
def complete_maintenance_plan(plan_id: int, payload: MaintenanceCompleteIn,
                              db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    plan = db.get(PreventiveMaintenancePlan, plan_id)
    if not plan: raise HTTPException(404, "Plan no existe")
    validate_technician_assignment(db, payload.technician.strip() or u.full_name or u.username)
    now_ = dt.datetime.now(dt.timezone.utc)
    previous_due = plan.next_due_at
    db.add(PreventiveMaintenanceLog(plan_id=plan.id, machine_id=plan.machine_id,
        due_at=previous_due, completed_at=now_, technician=payload.technician.strip() or u.full_name or u.username,
        notes=payload.notes.strip(), user_id=u.id))
    db.add(Intervention(machine_id=plan.machine_id, area_id=plan.machine.area_id,
        occurred_at=now_, technician=payload.technician.strip() or u.full_name or u.username,
        work_type="MANTENIMIENTO PREVENTIVO", task=plan.title,
        detail=payload.notes.strip(), island=plan.machine.island.number if plan.machine.island else "",
        jira_number="", pending=False, result="Mantenimiento preventivo ejecutado",
        notes=f"Plan preventivo #{plan.id} · frecuencia {plan.interval_days} días", created_by=u.id))
    plan.next_due_at = now_ + dt.timedelta(days=plan.interval_days)
    plan.notified_for_due_at = None
    audit(db, u, "COMPLETAR_MANTENIMIENTO", f"maintenance:{plan.id}",
        {"next_due_at": previous_due}, {"completed_at": now_, "next_due_at": plan.next_due_at})
    db.commit()
    return maintenance_plan_dict(plan)

@app.patch("/maintenance/plans/{plan_id}/active")
def set_maintenance_plan_active(plan_id: int, active: bool, db: Session = Depends(get_db), u=Depends(require("JEFE", "SUPERVISOR"))):
    plan = db.get(PreventiveMaintenancePlan, plan_id)
    if not plan: raise HTTPException(404, "Plan no existe")
    old = plan.active; plan.active = active
    audit(db, u, "CAMBIAR_ESTADO_PLAN_MANTENIMIENTO", f"maintenance:{plan.id}", {"active": old}, {"active": active})
    db.commit()
    return maintenance_plan_dict(plan)

@app.get("/maintenance/plans/{plan_id}/history")
def maintenance_plan_history(plan_id: int, db: Session = Depends(get_db), _=Depends(staff_user)):
    if not db.get(PreventiveMaintenancePlan, plan_id): raise HTTPException(404, "Plan no existe")
    rows = db.scalars(select(PreventiveMaintenanceLog).where(PreventiveMaintenanceLog.plan_id == plan_id)
                      .order_by(PreventiveMaintenanceLog.completed_at.desc())).all()
    return [{"id": row.id, "due_at": row.due_at, "completed_at": row.completed_at,
             "technician": row.technician, "notes": row.notes} for row in rows]

@app.get("/tickets/{ticket_id}/history")
def ticket_history(ticket_id: int, db: Session = Depends(get_db), u=Depends(current_user)):
    ticket = db.get(Ticket, ticket_id)
    if not ticket: raise HTTPException(404, "Ticket no existe")
    if u.role == "USUARIO" and ticket.created_by != u.id: raise HTTPException(404, "Ticket no existe")
    rows = db.scalars(select(TicketStatusHistory).where(TicketStatusHistory.ticket_id == ticket_id)
                      .order_by(TicketStatusHistory.at.desc())).all()
    return [{"at": r.at, "old_status": r.old_status, "new_status": r.new_status,
             "user_id": r.user_id, "note": r.note} for r in rows]

class TicketCommentIn(BaseModel):
    body: str

def ticket_comment_dict(comment: TicketComment):
    return {"id": comment.id, "ticket_id": comment.ticket_id, "at": comment.at,
            "body": comment.body, "author": comment.user.full_name or comment.user.username,
            "author_role": comment.user.role}

def accessible_ticket(ticket_id: int, user: User, db: Session) -> Ticket:
    ticket = db.get(Ticket, ticket_id)
    if not ticket or user.role == "USUARIO" and ticket.created_by != user.id:
        raise HTTPException(404, "Ticket no existe")
    return ticket

@app.get("/tickets/{ticket_id}/comments")
def list_ticket_comments(ticket_id: int, db: Session = Depends(get_db), u=Depends(current_user)):
    accessible_ticket(ticket_id, u, db)
    rows = db.scalars(select(TicketComment).where(TicketComment.ticket_id == ticket_id)
                      .order_by(TicketComment.at, TicketComment.id)).all()
    return [ticket_comment_dict(comment) for comment in rows]

@app.post("/tickets/{ticket_id}/comments", status_code=201)
def add_ticket_comment(ticket_id: int, payload: TicketCommentIn,
                       db: Session = Depends(get_db), u=Depends(current_user)):
    if u.role not in (*WRITERS, "ADMIN", "USUARIO"):
        raise HTTPException(403, "Sin permiso")
    ticket = accessible_ticket(ticket_id, u, db)
    body = payload.body.strip()
    if not body: raise HTTPException(422, "El mensaje no puede estar vacío")
    if len(body) > 4000: raise HTTPException(422, "El mensaje no puede superar 4000 caracteres")
    comment = TicketComment(ticket_id=ticket.id, user_id=u.id, body=body)
    db.add(comment)
    audit(db, u, "COMENTAR_TICKET", f"ticket:{ticket.id}", None, {"body": body})
    db.commit()
    db.refresh(comment)
    if u.role == "USUARIO":
        send_ticket_notification(f"Nuevo mensaje en ticket #{ticket.id}",
                                 f"{u.full_name or u.username} agregó información al ticket '{ticket.task}'.")
    else:
        requester = db.get(User, ticket.created_by)
        send_system_email([requester.username] if requester else [], f"Respuesta a tu ticket #{ticket.id}",
                          f"El equipo técnico respondió a '{ticket.task}':\n\n{body}")
    return ticket_comment_dict(comment)

class InterventionIn(BaseModel):
    occurred_at: dt.datetime | None = None
    ticket_id: int | None = None
    maintenance_plan_id: int | None = None
    machine_status: str | None = None
    machine_status_reason: str | None = None
    ticket_status: str | None = None
    ticket_status_note: str | None = None
    machine: str | None = None
    area: str | None = None
    technician: str | None = None
    shift: str | None = None
    work_type: str | None = None
    task: str
    detail: str = ""
    island: str | None = None
    jira_number: str | None = None
    pending: bool = False
    result: str | None = None
    notes: str | None = None
    parts_used: list[InventoryUseIn] = Field(default_factory=list, max_length=100)

def intervention_dict(i: Intervention):
    latest_follow_up = i.follow_ups[-1] if i.follow_ups else None
    pending = latest_follow_up.status != "RESUELTA" if latest_follow_up else i.pending
    return {"id": i.id, "occurred_at": i.occurred_at, "ticket_id": i.ticket_id,
            "machine": i.machine.number if i.machine_id and i.machine else None,
            "area": i.area.name if i.area_id and i.area else None, "technician": i.technician,
            "shift": i.shift, "work_type": i.work_type, "task": i.task, "detail": i.detail,
            "island": i.island, "jira_number": i.jira_number, "pending": pending,
            "follow_up_status": latest_follow_up.status if latest_follow_up else ("PENDIENTE" if i.pending else "RESUELTA"),
            "follow_up_note": latest_follow_up.note if latest_follow_up else None,
            "follow_up_by": (latest_follow_up.user.full_name or latest_follow_up.user.username) if latest_follow_up and latest_follow_up.user else None,
            "follow_up_at": latest_follow_up.at if latest_follow_up else None,
            "result": i.result, "notes": i.notes}

@app.post("/interventions", status_code=201)
def create_intervention(data: InterventionIn, db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    validate_technician_assignment(db, data.technician)
    if not data.task.strip(): raise HTTPException(422, "La tarea realizada es obligatoria")
    ticket = db.get(Ticket, data.ticket_id) if data.ticket_id else None
    if data.ticket_id and not ticket: raise HTTPException(422, "Ticket relacionado no existe")
    if data.ticket_status and not ticket:
        raise HTTPException(422, "Vincula un ticket para cambiar su estado")
    if data.ticket_status and data.ticket_status not in TICKET_STATES:
        raise HTTPException(422, "Estado de ticket inválido")
    if data.ticket_status_note and len(data.ticket_status_note) > 4000:
        raise HTTPException(422, "La nota del cambio no puede superar 4000 caracteres")
    machine = db.scalar(select(Machine).where(Machine.number == data.machine)) if data.machine else None
    if data.machine and not machine: raise HTTPException(422, "Máquina no existe")
    inventory_parts = lock_inventory_for_usage(data.parts_used, db)
    if ticket and ticket.machine_id:
        if machine and machine.id != ticket.machine_id: raise HTTPException(422, "La máquina no coincide con el ticket")
        machine = machine or db.get(Machine, ticket.machine_id)
    maintenance_plan = db.get(PreventiveMaintenancePlan, data.maintenance_plan_id) if data.maintenance_plan_id else None
    if data.maintenance_plan_id and not maintenance_plan:
        raise HTTPException(422, "Plan preventivo no existe")
    if maintenance_plan:
        if not data.work_type or "MANTENIMIENTO" not in data.work_type.upper():
            raise HTTPException(422, "Selecciona un tipo de trabajo de mantenimiento para vincular el plan")
        if not maintenance_plan.active:
            raise HTTPException(422, "El plan preventivo está inactivo")
        if machine and machine.id != maintenance_plan.machine_id:
            raise HTTPException(422, "El plan seleccionado corresponde a otra máquina")
        machine = machine or db.get(Machine, maintenance_plan.machine_id)
    if data.machine_status:
        if not machine: raise HTTPException(422, "Selecciona una máquina para actualizar su estado")
        if data.machine_status not in MACHINE_STATES: raise HTTPException(422, "Estado de máquina inválido")
        if not (data.machine_status_reason or "").strip(): raise HTTPException(422, "Indica el motivo del cambio de estado")
        if machine.status == data.machine_status: raise HTTPException(422, "La máquina ya tiene ese estado")
    area = db.scalar(select(Area).where(Area.name == data.area)) if data.area else (machine.area if machine else None)
    if data.area and not area: raise HTTPException(422, "Área no existe en el catálogo")
    occurred = data.occurred_at or dt.datetime.now(dt.timezone.utc)
    if occurred.tzinfo is None:
        occurred = occurred.replace(tzinfo=ZoneInfo("America/Santiago")).astimezone(dt.timezone.utc)
    if machine and data.machine_status:
        previous_status = machine.status
        if data.machine_status == "OPERATIVA":
            open_downtime = db.scalar(select(MachineStatusHistory).where(
                MachineStatusHistory.machine_id == machine.id,
                MachineStatusHistory.new_status == "FUERA_DE_SERVICIO",
                MachineStatusHistory.downtime_end.is_(None)).order_by(MachineStatusHistory.at.desc()))
            if open_downtime:
                open_downtime.downtime_end = occurred
                open_downtime.downtime_minutes = max(0, int((occurred - open_downtime.at).total_seconds() // 60))
        db.add(MachineStatusHistory(machine_id=machine.id, old_status=previous_status,
            new_status=data.machine_status, at=occurred, user_id=u.id,
            reason=data.machine_status_reason.strip(), ticket_id=data.ticket_id))
        machine.status = data.machine_status
        audit(db, u, "CAMBIO_ESTADO_DESDE_BITACORA", f"machine:{machine.number}", previous_status,
            {"status": machine.status, "reason": data.machine_status_reason.strip()})
    item = Intervention(occurred_at=occurred, ticket_id=data.ticket_id,
                        machine_id=machine.id if machine else None, area_id=area.id if area else None,
                        technician=data.technician.strip() if data.technician else None,
                        shift=data.shift, work_type=data.work_type, task=data.task.strip(), detail=data.detail,
                        island=(data.island or "").strip(), jira_number=(data.jira_number or "").strip(),
                        pending=data.pending, result=data.result, notes=data.notes, created_by=u.id)
    db.add(item); db.flush()
    low_stock_notifications = record_inventory_usage(
        data.parts_used, inventory_parts, db, u, ticket_id=data.ticket_id,
        machine_id=machine.id if machine else None,
        technician=(data.technician or "").strip() or None,
        reference=f"Consumo en bitácora · Intervención #{item.id} · {data.task.strip()}",
        audit_action="CONSUMO_INVENTARIO_DESDE_BITACORA")
    if maintenance_plan:
        previous_due = maintenance_plan.next_due_at
        db.add(PreventiveMaintenanceLog(plan_id=maintenance_plan.id, machine_id=maintenance_plan.machine_id,
            due_at=previous_due, completed_at=occurred,
            technician=(data.technician or "").strip() or u.full_name or u.username,
            notes=(data.detail or "").strip() or (data.notes or "").strip(), user_id=u.id))
        maintenance_plan.next_due_at = occurred + dt.timedelta(days=maintenance_plan.interval_days)
        maintenance_plan.notified_for_due_at = None
        audit(db, u, "ACTUALIZAR_PLAN_DESDE_BITACORA", f"maintenance:{maintenance_plan.id}",
            {"next_due_at": previous_due}, {"completed_at": occurred, "next_due_at": maintenance_plan.next_due_at,
             "intervention_id": item.id})
    if ticket and data.ticket_status and data.ticket_status != ticket.status:
        old_status = ticket.status
        ticket.status = data.ticket_status
        note = data.ticket_status_note.strip() if data.ticket_status_note else ""
        note = note or f"Estado actualizado al registrar intervención #{item.id}"
        db.add(TicketStatusHistory(ticket_id=ticket.id, old_status=old_status,
                                   new_status=ticket.status, user_id=u.id, note=note))
        audit(db, u, "ACTUALIZAR_TICKET_DESDE_BITACORA", f"ticket:{ticket.id}",
              {"status": old_status}, {"status": ticket.status, "note": note,
                                        "intervention_id": item.id})
    audit(db, u, "REGISTRAR_INTERVENCION", f"intervention:{item.id}", None, data.model_dump())
    db.commit()
    notify_inventory_below_minimum(low_stock_notifications)
    return intervention_dict(item)

@app.get("/interventions")
def list_interventions(search_text: str | None = None, machine: str | None = None, technician: str | None = None,
                       date_from: dt.date | None = None, date_to: dt.date | None = None,
                       area: str | None = None, shift: str | None = None,
                       work_type: str | None = None, pending_only: bool = False,
                       page: int = Query(1, ge=1), size: int = Query(50, le=200),
                       db: Session = Depends(get_db), _=Depends(staff_user)):
    query = select(Intervention).options(selectinload(Intervention.follow_ups).selectinload(InterventionFollowUp.user))
    if search_text:
        term = f"%{search_text.strip()}%"
        query = query.where((Intervention.task.ilike(term)) | (Intervention.detail.ilike(term)) |
                            (Intervention.technician.ilike(term)) | (Intervention.jira_number.ilike(term)))
    local_tz = ZoneInfo("America/Santiago")
    if date_from:
        start_at = dt.datetime.combine(date_from, dt.time.min, tzinfo=local_tz).astimezone(dt.timezone.utc)
        query = query.where(Intervention.occurred_at >= start_at)
    if date_to:
        end_at = dt.datetime.combine(date_to + dt.timedelta(days=1), dt.time.min, tzinfo=local_tz).astimezone(dt.timezone.utc)
        query = query.where(Intervention.occurred_at < end_at)
    if machine:
        machine_row = db.scalar(select(Machine).where(Machine.number == machine))
        query = query.where(Intervention.machine_id == machine_row.id) if machine_row else query.where(Intervention.id == -1)
    if technician: query = query.where(Intervention.technician.ilike(f"%{technician.strip()}%"))
    if area: query = query.where(Intervention.area.has(Area.name.ilike(f"%{area.strip()}%")))
    if shift: query = query.where(Intervention.shift == shift)
    if work_type: query = query.where(Intervention.work_type.ilike(f"%{work_type.strip()}%"))
    if pending_only:
        latest_followup = (select(InterventionFollowUp.intervention_id.label("intervention_id"),
                                  func.max(InterventionFollowUp.id).label("followup_id"))
                           .group_by(InterventionFollowUp.intervention_id).subquery())
        query = query.outerjoin(latest_followup, latest_followup.c.intervention_id == Intervention.id)
        query = query.outerjoin(InterventionFollowUp, InterventionFollowUp.id == latest_followup.c.followup_id)
        query = query.where(or_(and_(latest_followup.c.followup_id.is_(None), Intervention.pending.is_(True)),
                                InterventionFollowUp.status.in_({"PENDIENTE", "RECIBIDA"})))
    total = db.scalar(select(func.count()).select_from(query.subquery()))
    rows = db.scalars(query.order_by(Intervention.occurred_at.desc()).offset((page-1)*size).limit(size)).all()
    return {"total": total, "items": [intervention_dict(i) for i in rows]}

class InterventionFollowUpIn(BaseModel):
    status: str
    note: str = Field(default="", max_length=4000)

@app.get("/interventions/pending")
def list_pending_interventions(page: int = Query(1, ge=1), size: int = Query(8, ge=1, le=100),
                               db: Session = Depends(get_db), _=Depends(staff_user)):
    latest_followup = (select(InterventionFollowUp.intervention_id.label("intervention_id"),
                              func.max(InterventionFollowUp.id).label("followup_id"))
                       .group_by(InterventionFollowUp.intervention_id).subquery())
    query = (select(Intervention).options(selectinload(Intervention.follow_ups).selectinload(InterventionFollowUp.user))
             .outerjoin(latest_followup, latest_followup.c.intervention_id == Intervention.id)
             .outerjoin(InterventionFollowUp, InterventionFollowUp.id == latest_followup.c.followup_id)
             .where(or_(and_(latest_followup.c.followup_id.is_(None), Intervention.pending.is_(True)),
                        InterventionFollowUp.status.in_({"PENDIENTE", "RECIBIDA"}))))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    received_total = db.scalar(select(func.count()).select_from(query.where(
        InterventionFollowUp.status == "RECIBIDA").subquery())) or 0
    rows = db.scalars(query.order_by(Intervention.occurred_at.desc()).offset((page-1)*size).limit(size)).all()
    return {"total": total, "received_total": received_total,
            "items": [intervention_dict(row) for row in rows]}

@app.post("/interventions/{intervention_id}/follow-ups", status_code=201)
def add_intervention_follow_up(intervention_id: int, payload: InterventionFollowUpIn,
                              db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    if payload.status not in {"PENDIENTE", "RECIBIDA", "RESUELTA"}:
        raise HTTPException(422, "El seguimiento debe quedar pendiente, recibido o resuelto")
    intervention = db.get(Intervention, intervention_id)
    if not intervention: raise HTTPException(404, "Intervención no existe")
    previous = intervention_dict(intervention)
    follow_up = InterventionFollowUp(intervention_id=intervention.id, status=payload.status,
                                     note=payload.note.strip(), user_id=u.id)
    db.add(follow_up)
    audit(db, u, "SEGUIMIENTO_INTERVENCION", f"intervention:{intervention.id}",
          {"pending": previous["pending"]}, {"status": payload.status, "note": payload.note.strip()})
    db.commit()
    db.refresh(intervention)
    return intervention_dict(intervention)

class DailyReportIn(BaseModel):
    report_date: dt.date
    shift: str | None = None

def report_recipients(env_key: str) -> list[str]:
    raw = email_recipient_settings().get(env_key, "")
    return [address.strip() for address in re.split(r"[;,]", raw) if address.strip()]

def daily_report_data(report_date: dt.date, shift: str | None, db: Session):
    local_tz = ZoneInfo("America/Santiago")
    start_at = dt.datetime.combine(report_date, dt.time.min, tzinfo=local_tz).astimezone(dt.timezone.utc)
    end_at = dt.datetime.combine(report_date + dt.timedelta(days=1), dt.time.min, tzinfo=local_tz).astimezone(dt.timezone.utc)
    query = select(Intervention).options(selectinload(Intervention.follow_ups).selectinload(InterventionFollowUp.user)).where(Intervention.occurred_at >= start_at,
                                       Intervention.occurred_at < end_at)
    if shift: query = query.where(Intervention.shift == shift)
    rows = db.scalars(query.order_by(Intervention.occurred_at)).all()
    items = [intervention_dict(row) for row in rows]
    machine_status_counts = dict(db.execute(select(Machine.status, func.count()).group_by(Machine.status)).all())
    machine_total = sum(machine_status_counts.values())
    machine_status_labels = {
        "OPERATIVA": "Operativa", "OPERATIVA_CON_OBSERVACION": "Operativa con observación",
        "EN_MANTENIMIENTO": "En mantenimiento", "FUERA_DE_SERVICIO": "Fuera de servicio",
        "PENDIENTE_DE_REPUESTO": "Pendiente de repuesto", "RETIRADA": "Retirada",
    }
    machine_fleet = {
        "total": machine_total,
        "operating": machine_status_counts.get("OPERATIVA", 0),
        "operating_with_observation": machine_status_counts.get("OPERATIVA_CON_OBSERVACION", 0),
        "out_of_service": machine_status_counts.get("FUERA_DE_SERVICIO", 0),
        "in_maintenance": machine_status_counts.get("EN_MANTENIMIENTO", 0),
        "waiting_parts": machine_status_counts.get("PENDIENTE_DE_REPUESTO", 0),
        "retired": machine_status_counts.get("RETIRADA", 0),
        "worked_on": len({row.machine_id for row in rows if row.machine_id is not None}),
        "by_status": [{"status": status, "label": machine_status_labels.get(status, status),
                       "count": count} for status, count in sorted(machine_status_counts.items())],
    }
    # Use the same date/shift selection as the daily activities, and their current follow-up state.
    handoff_items = [item for item in items if item["pending"]]
    handoff_total = len(handoff_items)
    return {
        "date": report_date.isoformat(), "shift": shift or "Todos los turnos",
        "total": len(items), "pending_total": sum(1 for item in items if item["pending"]),
        "received_total": sum(1 for item in items if item["pending"] and item["follow_up_status"] == "RECIBIDA"),
        "open_handoff_total": handoff_total,
        "handoff_received_total": sum(1 for item in handoff_items if item["follow_up_status"] == "RECIBIDA"),
        "handoff_items": handoff_items,
        "machine_fleet": machine_fleet,
        "items": items,
    }

def daily_report_html(report: dict, sender: str) -> str:
    esc = lambda value: html_lib.escape(str(value or ""), quote=True)
    handoff_rows = []
    for item in report["handoff_items"]:
        received = item["follow_up_status"] == "RECIBIDA"
        handoff_rows.append("<tr>" + "".join(f"<td>{value}</td>" for value in (
            esc(item["machine"] or "—"), esc(item["task"]),
            esc("Recibida" if received else "Pendiente de recepción"),
            esc(item["follow_up_by"] if received else "—"), esc(item["follow_up_note"] or item["detail"] or "—"),
        )) + "</tr>")
    if not handoff_rows:
        handoff_rows.append('<tr><td colspan="5" style="text-align:center;color:#6b7280">No hay actividades pendientes para la fecha y turno seleccionados.</td></tr>')
    machine_rows = ["<tr>" + "".join(f"<td>{value}</td>" for value in (
        esc(item["label"]), str(item["count"]))) + "</tr>" for item in report["machine_fleet"]["by_status"]]
    if not machine_rows:
        machine_rows.append('<tr><td colspan="2" style="text-align:center;color:#6b7280">Sin máquinas registradas.</td></tr>')
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>Bitácora diaria</title></head>
    <body style="margin:0;background:#f3f6fa;font-family:Arial,sans-serif;color:#243247"><div style="max-width:1100px;margin:24px auto;background:#fff;border-radius:10px;overflow:hidden">
    <header style="padding:28px 32px;background:#205ca8;color:#fff"><div style="font-size:12px;letter-spacing:1px">GESTIÓN TÉCNICA · CASINO &amp; RESORT</div><h1 style="font-size:24px;margin:10px 0 4px">Informe diario de actividades</h1><div>{esc(report['date'])} · {esc(report['shift'])}</div></header>
    <section style="padding:24px 32px"><div style="display:flex;gap:24px;margin-bottom:22px"><div><b style="font-size:24px">{report['total']}</b><br><span>Actividades</span></div><div><b style="font-size:24px;color:#d27b1e">{report['pending_total']}</b><br><span>Pendientes</span></div><div><b style="font-size:24px;color:#278756">{report['received_total']}</b><br><span>Recibidas por el turno entrante</span></div><div><b>Enviado por</b><br><span>{esc(sender)}</span></div></div>
    <h2 style="font-size:18px;margin:8px 0 4px">Estado actual del parque · {report['machine_fleet']['total']} máquinas · {report['machine_fleet']['worked_on']} intervenidas en el período</h2><p style="font-size:12px;color:#6b7280;margin:0 0 10px">Estado operativo al momento de generar el informe</p><div style="overflow:auto"><table style="width:100%;border-collapse:collapse;font-size:13px"><thead><tr style="background:#eef3fa"><th style="text-align:left;padding:10px;border-bottom:1px solid #dfe6ef">Estado</th><th style="text-align:right;padding:10px;border-bottom:1px solid #dfe6ef">Máquinas</th></tr></thead><tbody>{''.join(machine_rows)}</tbody></table></div>
    <h2 style="font-size:18px;margin:28px 0 4px">Actividades pendientes del día</h2><p style="font-size:12px;color:#6b7280;margin:0 0 10px">{report['open_handoff_total']} pendientes · {report['handoff_received_total']} recibidas por el turno entrante</p><div style="overflow:auto"><table style="width:100%;border-collapse:collapse;font-size:13px"><thead><tr style="background:#eef3fa">{''.join(f'<th style="text-align:left;padding:10px;border-bottom:1px solid #dfe6ef">{label}</th>' for label in ('Máquina','Intervención','Estado','Recibida por','Nota de seguimiento'))}</tr></thead><tbody>{''.join(handoff_rows)}</tbody></table></div></section>
    </div></body></html>"""

@app.get("/reports/daily")
def preview_daily_report(report_date: dt.date | None = None, shift: str | None = None,
                         db: Session = Depends(get_db), _=Depends(staff_user)):
    report_date = report_date or dt.date.today()
    email_config = active_email_settings()
    return {**daily_report_data(report_date, shift, db),
            "recipients": report_recipients("REPORT_EMAIL_TO"),
            "cc": report_recipients("REPORT_EMAIL_CC"),
            "email_configured": bool(email_config["smtp_host"] and email_config["smtp_user"]
                                     and email_config["smtp_password"]
                                     and report_recipients("REPORT_EMAIL_TO"))}

@app.post("/reports/daily/email")
def send_daily_report(payload: DailyReportIn, db: Session = Depends(get_db), u=Depends(require(*WRITERS))):
    config = active_email_settings()
    recipients = report_recipients("REPORT_EMAIL_TO")
    cc = report_recipients("REPORT_EMAIL_CC")
    if not config["smtp_host"] or not config["smtp_user"] or not config["smtp_password"]:
        raise HTTPException(503, "Configura el servidor SMTP en Administración → Configuración")
    if not recipients:
        raise HTTPException(503, "Configura al menos un destinatario de informe en Administración → Configuración")
    report = daily_report_data(payload.report_date, payload.shift, db)
    sender_name = u.full_name or u.username
    subject = f"Bitácora diaria · {payload.report_date:%d/%m/%Y} · {payload.shift or 'Todos los turnos'}"
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr((config["smtp_from_name"], config["smtp_from"] or config["smtp_user"]))
    message["To"] = ", ".join(recipients)
    if cc: message["Cc"] = ", ".join(cc)
    message.set_content(f"Informe de bitácora del {report['date']} ({report['shift']}). Actividades: {report['total']}; pendientes: {report['pending_total']}. Enviado por {sender_name}.")
    message.add_alternative(daily_report_html(report, sender_name), subtype="html")
    if not send_email_message(message, config):
        raise HTTPException(502, "No se pudo enviar el correo; revisa la configuración SMTP y los registros del backend")
    audit(db, u, "ENVIAR_INFORME_BITACORA", "daily_report",
          None, {"date": report["date"], "shift": report["shift"], "to": recipients, "cc": cc})
    db.commit()
    return {"sent": True, "recipients": recipients, "cc": cc, "total": report["total"]}

@app.get("/kpi/summary")
def kpi(db: Session = Depends(get_db), _=Depends(staff_user)):
    by_status = dict(db.execute(select(Machine.status, func.count()).group_by(Machine.status)).all())
    mttr = db.scalar(select(func.avg(MachineStatusHistory.downtime_minutes)))
    ticket_statuses = dict(db.execute(select(Ticket.status, func.count()).group_by(Ticket.status)).all())
    period_start = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=30)
    unresolved = case((Ticket.status.in_({"RESUELTO", "CERRADO", "CANCELADO"}), 0), else_=1)
    now_ = dt.datetime.now(dt.timezone.utc)
    open_ticket_filter = Ticket.status.notin_({"RESUELTO", "CERRADO", "CANCELADO"})
    tickets_unassigned = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, (Ticket.technician.is_(None)) | (func.trim(Ticket.technician) == ""))) or 0
    tickets_over_24h = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.created_at <= now_ - dt.timedelta(hours=24))) or 0
    tickets_over_72h = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.created_at <= now_ - dt.timedelta(hours=72))) or 0
    tickets_sla_overdue = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.sla_due_at.is_not(None), Ticket.sla_due_at <= now_)) or 0
    tickets_sla_approaching = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.sla_due_at > now_,
        Ticket.sla_due_at <= now_ + dt.timedelta(hours=24))) or 0
    maintenance_overdue = db.scalar(select(func.count(PreventiveMaintenancePlan.id)).where(
        PreventiveMaintenancePlan.active.is_(True), PreventiveMaintenancePlan.next_due_at < now_)) or 0
    maintenance_upcoming = db.scalar(select(func.count(PreventiveMaintenancePlan.id)).where(
        PreventiveMaintenancePlan.active.is_(True), PreventiveMaintenancePlan.next_due_at >= now_,
        PreventiveMaintenancePlan.next_due_at <= now_ + dt.timedelta(days=7))) or 0
    due_maintenance_plans = db.scalars(select(PreventiveMaintenancePlan).where(
        PreventiveMaintenancePlan.active.is_(True),
        PreventiveMaintenancePlan.next_due_at <= now_ + dt.timedelta(days=7))
        .order_by(PreventiveMaintenancePlan.next_due_at).limit(5)).all()
    tickets_urgent = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.priority.in_({"CRÍTICA", "ALTA"}))) or 0
    tickets_under_24h = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.created_at > now_ - dt.timedelta(hours=24))) or 0
    tickets_24_to_72h = db.scalar(select(func.count(Ticket.id)).where(
        open_ticket_filter, Ticket.created_at <= now_ - dt.timedelta(hours=24),
        Ticket.created_at > now_ - dt.timedelta(hours=72))) or 0
    local_tz = ZoneInfo("America/Santiago")
    local_today = dt.datetime.now(local_tz).date()
    trend_start_date = local_today - dt.timedelta(days=6)
    trend_start = dt.datetime.combine(trend_start_date, dt.time.min, tzinfo=local_tz).astimezone(dt.timezone.utc)
    trend_end = dt.datetime.combine(local_today + dt.timedelta(days=1), dt.time.min,
                                    tzinfo=local_tz).astimezone(dt.timezone.utc)
    trend_days = {(trend_start_date + dt.timedelta(days=offset)).isoformat(): {"created": 0, "resolved": 0}
                  for offset in range(7)}
    created_rows = db.scalars(select(Ticket.created_at).where(
        Ticket.created_at >= trend_start, Ticket.created_at < trend_end)).all()
    for created_at in created_rows:
        created_utc = created_at.replace(tzinfo=dt.timezone.utc) if created_at.tzinfo is None else created_at
        day = created_utc.astimezone(local_tz).date().isoformat()
        if day in trend_days:
            trend_days[day]["created"] += 1
    resolved_rows = db.execute(select(TicketStatusHistory.at, TicketStatusHistory.old_status,
                                      TicketStatusHistory.new_status).where(
        TicketStatusHistory.at >= trend_start, TicketStatusHistory.at < trend_end,
        TicketStatusHistory.new_status.in_({"RESUELTO", "CERRADO"}))).all()
    for resolved_at, old_status, _new_status in resolved_rows:
        if _new_status == "CERRADO" and old_status == "RESUELTO":
            continue
        resolved_utc = resolved_at.replace(tzinfo=dt.timezone.utc) if resolved_at.tzinfo is None else resolved_at
        day = resolved_utc.astimezone(local_tz).date().isoformat()
        if day in trend_days:
            trend_days[day]["resolved"] += 1
    first_resolution = (select(TicketStatusHistory.ticket_id.label("ticket_id"),
                               func.min(TicketStatusHistory.at).label("resolved_at"))
                        .where(TicketStatusHistory.new_status.in_({"RESUELTO", "CERRADO"}))
                        .group_by(TicketStatusHistory.ticket_id).subquery())
    sla_completed = (select(Ticket.id.label("ticket_id"),
                            case((first_resolution.c.resolved_at <= Ticket.sla_due_at, 1), else_=0).label("met_sla"))
                     .join(first_resolution, first_resolution.c.ticket_id == Ticket.id)
                     .where(Ticket.sla_due_at.is_not(None),
                            Ticket.created_at >= period_start,
                            Ticket.status.in_({"RESUELTO", "CERRADO"})).subquery())
    sla_completed_count = db.scalar(select(func.count()).select_from(sla_completed)) or 0
    sla_on_time_count = db.scalar(select(func.coalesce(func.sum(sla_completed.c.met_sla), 0))
                                   .select_from(sla_completed)) or 0
    machine_rows = db.execute(
        select(Machine.number, func.count(Ticket.id), func.sum(unresolved), func.max(Ticket.created_at))
        .join(Ticket, Ticket.machine_id == Machine.id)
        .where(Ticket.created_at >= period_start)
        .group_by(Machine.id, Machine.number)
        .order_by(func.count(Ticket.id).desc(), func.sum(unresolved).desc(), Machine.number)
        .limit(5)
    ).all()
    technician_name = func.trim(Intervention.technician)
    latest_intervention_followup = (select(InterventionFollowUp.intervention_id.label("intervention_id"),
                                          func.max(InterventionFollowUp.id).label("followup_id"))
                                   .group_by(InterventionFollowUp.intervention_id).subquery())
    intervention_pending = case(
        (latest_intervention_followup.c.followup_id.is_(None), Intervention.pending),
        (InterventionFollowUp.status.in_({"PENDIENTE", "RECIBIDA"}), True), else_=False)
    technician_rows = db.execute(
        select(technician_name, func.count(Intervention.id),
               func.sum(case((intervention_pending, 1), else_=0)),
               func.max(Intervention.occurred_at))
        .outerjoin(latest_intervention_followup,
                   latest_intervention_followup.c.intervention_id == Intervention.id)
        .outerjoin(InterventionFollowUp,
                   InterventionFollowUp.id == latest_intervention_followup.c.followup_id)
        .where(Intervention.occurred_at >= period_start,
               Intervention.technician.is_not(None), technician_name != "")
        .group_by(technician_name)
        .order_by(func.count(Intervention.id).desc(), technician_name)
        .limit(5)
    ).all()
    pending_ticket_rows = db.scalars(
        select(Ticket)
        .where(Ticket.status.notin_({"RESUELTO", "CERRADO", "CANCELADO"}))
        .order_by(case((Ticket.sla_due_at.is_not(None) & (Ticket.sla_due_at <= now_), 0), else_=1),
                       Ticket.sla_due_at.asc().nullslast(),
                       case((Ticket.priority == "CRÍTICA", 0), (Ticket.priority == "ALTA", 1),
                       (Ticket.priority == "NORMAL", 2), else_=3), Ticket.created_at.asc())
        .limit(6)
    ).all()
    return {"tickets_total": db.scalar(select(func.count(Ticket.id))),
            "tickets_by_status": ticket_statuses,
            "machines_by_status": by_status,
            "ticket_attention_counts": {"unassigned": tickets_unassigned,
                                        "over_24h": tickets_over_24h,
                                        "over_72h": tickets_over_72h,
                                        "under_24h": tickets_under_24h,
                                        "24_to_72h": tickets_24_to_72h,
                                        "sla_overdue": tickets_sla_overdue,
                                        "sla_approaching": tickets_sla_approaching,
                                        "urgent": tickets_urgent},
            "maintenance_attention": {"overdue": maintenance_overdue,
                "upcoming_7d": maintenance_upcoming,
                "plans": [{"id": plan.id, "machine_number": plan.machine.number,
                           "title": plan.title, "next_due_at": plan.next_due_at,
                           "overdue": plan.next_due_at < now_} for plan in due_maintenance_plans]},
            "ticket_age_buckets": {"under_24h": tickets_under_24h,
                                   "24_to_72h": tickets_24_to_72h,
                                   "over_72h": tickets_over_72h},
            "ticket_trend_7d": [{"date": day, **counts} for day, counts in trend_days.items()],
            "sla_performance_30d": {"resolved_with_sla": sla_completed_count,
                                     "resolved_on_time": sla_on_time_count,
                                     "compliance_percent": round(sla_on_time_count / sla_completed_count * 100, 1)
                                     if sla_completed_count else None},
            "mttr_minutes": round(float(mttr), 1) if mttr is not None else None,
            "period_days": 30,
            "top_problem_machines": [{"number": row[0], "ticket_count": row[1],
                                      "open_tickets": row[2] or 0, "last_ticket_at": row[3]}
                                     for row in machine_rows],
            "top_technicians": [{"name": row[0], "interventions": row[1],
                                 "pending_interventions": row[2] or 0, "last_activity_at": row[3]}
                                for row in technician_rows],
            "attention_tickets": [ticket_dict(ticket) for ticket in pending_ticket_rows]}
